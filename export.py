"""The school spreadsheet: every class's records in one Excel workbook, for admins.

/admin/export describes the workbook and /admin/export.xlsx downloads it. Each area of the records has its own sheet,
one row per record, with a styled header, a frozen top row, filters and fitted column widths, so it can be sorted,
filtered and pivoted straight away. Times are London time.
"""
import io
import json
import re
from datetime import date, datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from flask import Blueprint, Response, current_app, g
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

import auth
from tv import natural

ROOT = Path(__file__).resolve().parent
LONDON = ZoneInfo("Europe/London")
HOUSES = ("Makkah", "Al-Aqsa", "Madinah")
TOKENS = (("tg", "Green token", 5), ("tb", "Blue token", 20), ("tr", "Red token", 50))  # TOK in src/page.html
# What a token was given for, when it was given on the Record page's House points (HPR in src/page.html).
REASONS = {"bk": "Book completion", "lv": "Qaidah or Memorisation level", "hw": "Additional class or homework",
           "ec": "Extra curricular award", "cp": "Exceptional class participation"}
BOOKS = {"q": "Qaidah Book", "t": "Tajweed Book"}  # BOOKS in src/page.html
QAIDAH_BOOK = 13  # Quran levels 1 to 13 are the Qaidah Book, the rest the Tajweed Book (QBOOK in src/page.html)
WARNINGS = ("None", "Warning 1", "Warning 2", "Miss break")  # WL in src/page.html
NAVY, WHITE, BAND, LINE, MUTED = "003B6B", "FFFFFF", "F3F5F8", "C9D1DB", "6B7280"
DAY, WHEN = "d mmm yyyy", "d mmm yyyy hh:mm"
MIME = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"

bp = Blueprint("export", __name__)

# Each sheet: its title, what it holds, and its columns (header, least width, number format).
SHEETS = {
    "students": ("Students", "One row per student: their details, where they are and their totals.", [
        ("Class", 22, None), ("Student ID", 11, None), ("Student", 22, None), ("Family ID", 10, None),
        ("Date of birth", 13, DAY), ("Gender", 9, None), ("Safar year", 10, None), ("School year", 11, None), ("Islamic Studies teacher", 22, None),
        ("House", 10, None), ("Memo target", 16, None), ("Memo passes", 9, "0"),
        ("Passes this month", 9, "0"), ("Last memo pass", 13, DAY), ("Memo fails", 9, "0"),
        ("Qaidah level", 9, "0"), ("Quran days scored", 9, "0"), ("Quran average", 9, "0.0"),
        ("Merits", 9, "0"), ("Warning", 11, None), ("House points", 9, "0"), ("Green tokens", 9, "0"),
        ("Blue tokens", 9, "0"), ("Red tokens", 9, "0"), ("Stars of the week", 9, "0")]),
    "classes": ("Classes", "One row per class, with its teachers and class totals.", [
        ("Class", 22, None), ("Teachers", 26, None), ("Students", 9, "0"), ("First class day", 13, DAY),
        ("Memo passes", 9, "0"), ("Passes this month", 9, "0"), ("Memo fails", 9, "0"),
        ("Level tests", 9, "0"), ("Quran average", 9, "0.0"), ("Merits", 9, "0"),
        ("With a warning", 9, "0"), ("House points", 9, "0")]),
    "quran": ("Quran scores", "Every daily Quran score, 0 to 5 or Absent, one row per student per day.", [
        ("Date", 13, DAY), ("Class", 22, None), ("Student ID", 11, None), ("Student", 22, None),
        ("Score", 8, "0")]),
    "memo": ("Memo passes and fails", "Every Memo pass and fail, newest first.", [
        ("Date and time", 17, WHEN), ("Class", 22, None), ("Student ID", 11, None), ("Student", 22, None),
        ("Result", 11, None), ("Target", 18, None)]),
    "tests": ("Level tests", "End-of-level tests and book passes recorded on the Record page.", [
        ("Date", 13, DAY), ("Class", 22, None), ("Student ID", 11, None), ("Student", 22, None),
        ("Test", 8, None), ("Level or target", 18, None), ("Score out of 5", 9, "0"), ("Result", 11, None),
        ("Blue token", 9, None)]),
    "rewards": ("Rewards", "Every merit and house token given, with its house points.", [
        ("Date and time", 17, WHEN), ("Class", 22, None), ("Student ID", 11, None), ("Student", 22, None),
        ("House", 10, None), ("Award", 30, None), ("House points", 9, "0")]),
    "stars": ("Stars of the week", "Each week's stars, by the Monday the week starts.", [
        ("Week starting", 13, DAY), ("Class", 22, None), ("Student ID", 11, None), ("Student", 22, None),
        ("House", 10, None)]),
    "warnings": ("Warnings", "Students on a warning or a missed break right now.", [
        ("Class", 22, None), ("Student ID", 11, None), ("Student", 22, None), ("House", 10, None),
        ("Warning", 12, None)]),
    "staff": ("Staff", "Staff accounts and the classes each one can open.", [
        ("Name", 22, None), ("Email", 28, None), ("Role", 13, None), ("Classes", 36, None),
        ("Status", 11, None), ("Added", 13, DAY)]),
}


def _juz():
    """The numbered Juz items after the book, read from the dashboard itself so the names never drift."""
    page = (ROOT / "Master_Dashboard_standalone.html").read_text(encoding="utf-8")
    found = re.search(r"var JUZ=(\[\[.*?\]\]);", page, re.S)
    items = []
    for _, first, surahs in json.loads(found.group(1)) if found else []:
        items += [f"{first + k} {name}" for k, name in enumerate(surahs.split(","))]
    return items + ["Quran complete"]


JUZ_ITEMS = _juz()


def quran_level(n):
    return f"{'Tajweed' if n > QAIDAH_BOOK else 'Qaidah'} Level {n}"


def target_label(n, levels):
    """A Memo target number as the dashboard shows it: 'L3 T5' in the book, then '205 Al-Buruj' and so on."""
    for level, size in enumerate(levels, 1):
        if n <= size:
            return f"L{level} T{max(n, 1)}"
        n -= size
    return JUZ_ITEMS[min(n, len(JUZ_ITEMS)) - 1] if JUZ_ITEMS else f"Item {n}"


def when(ms):
    """A browser timestamp (milliseconds) as a London date and time, which Excel stores without a time zone."""
    return datetime.fromtimestamp(ms / 1000, LONDON).replace(tzinfo=None, microsecond=0)


def day(text):
    return date.fromisoformat(text)


def birth(text):
    """A date of birth from the student's profile, or nothing when it isn't recorded."""
    return day(text) if text else None


class Sheet:
    """One worksheet: its columns (header, width hint, number format) and rows."""

    def __init__(self, title, about, columns):
        self.title, self.about, self.columns, self.rows = title, about, columns, []

    def add(self, *row):
        self.rows.append(row)


def gather(db):
    """Reads every class, student and staff account into the workbook's sheets."""
    teachers = {}
    for row in db.execute("SELECT m.class_id, u.name FROM classroom_members m JOIN users u ON u.id = m.user_id "
                          "ORDER BY u.name").fetchall():
        teachers.setdefault(row["class_id"], []).append(row["name"])
    classes = [(row["id"], json.loads(row["json"])) for row in db.execute("SELECT id, json FROM classrooms").fetchall()]
    classes.sort(key=lambda c: natural(c[1].get("name", c[0])))
    users = db.execute("SELECT u.id, u.name, u.email, u.role, u.disabled, u.password_hash, u.created_at FROM users u "
                       "ORDER BY u.role, u.name").fetchall()
    member_of = {}
    for row in db.execute("SELECT class_id, user_id FROM classroom_members").fetchall():
        member_of.setdefault(row["user_id"], []).append(row["class_id"])
    names = {cid: state.get("name", cid) for cid, state in classes}

    now = datetime.now(LONDON)
    month = datetime(now.year, now.month, 1, tzinfo=LONDON).timestamp() * 1000
    sheets = {key: Sheet(*spec) for key, spec in SHEETS.items()}
    houses = dict.fromkeys(HOUSES, 0)
    memo, rewards = [], []
    for cid, state in classes:
        cname, levels = names[cid], state.get("L") or [1]
        students = sorted(state.get("S", []), key=lambda s: natural(s.get("name", "")))
        totals = {"passes": 0, "month": 0, "fails": 0, "tests": 0, "merits": 0, "warned": 0, "points": 0, "q": []}
        for s in students:
            sid, name, house = s.get("sid", ""), s.get("name", ""), s.get("house", "")
            log, fails, merits = s.get("log", []), s.get("fl", []), s.get("ml", [])
            # -1 is absent: listed on the Quran sheet, left out of every average.
            scores = s.get("q", {})
            marks = [v for v in scores.values() if v >= 0]
            points = sum(len(s.get(k, [])) * v for k, _, v in TOKENS)
            started = s.get("target", 1) > 1 or log or fails
            sheets["students"].add(
                cname, sid, name, s.get("fid", ""), birth(s.get("dob", "")), s.get("gender", ""), s.get("year", ""), s.get("sy", ""), s.get("ist", ""), house,
                target_label(s.get("target", 1), levels) if started else "No data", s.get("passes", 0),
                sum(1 for t in log if t >= month), when(max(log)).date() if log else None, len(fails),
                s.get("qa") or None, len(marks), round(sum(marks) / len(marks), 2) if marks else None,
                s.get("merits", 0), WARNINGS[s.get("warn", 0)], points, *(len(s.get(k, [])) for k, _, _ in TOKENS),
                len(s.get("sw", [])))
            for d, score in sorted(scores.items()):
                sheets["quran"].add(day(d), cname, sid, name, "Absent" if score < 0 else score)
            labels = s.get("lx", [])
            for k, t in enumerate(log):
                lt = s.get("lt", [])
                label = (labels[k] if k < len(labels) else "") or (target_label(lt[k], levels) if k < len(lt) else "")
                memo.append((when(t), cname, sid, name, "Pass", label))
            memo += [(when(t), cname, sid, name, "Not passed", "") for t in fails]
            for test in s.get("tests", []):
                sheets["tests"].add(day(test["d"]), cname, sid, name, "Quran" if test["k"] == "q" else "Memo",
                                    quran_level(test["n"]) if test["k"] == "q" else target_label(test["n"], levels),
                                    test.get("sc"), "Passed" if test["ok"] else "Not passed", "Yes" if test.get("tk") else "")
            for b in s.get("bp", []):
                sheets["tests"].add(when(b["t"]).date(), cname, sid, name, "Book", BOOKS.get(b["b"], "Book"), None, "Passed", "")
            rewards += [(when(t), cname, sid, name, house, "Merit", None) for t in merits]
            why = {x["t"]: REASONS.get(x["r"]) for x in s.get("hp", []) if isinstance(x, dict)}
            for key, award, value in TOKENS:
                rewards += [(when(t), cname, sid, name, house, f"{award}: {why[t]}" if why.get(t) else award, value)
                            for t in s.get(key, [])]
            for monday in sorted(s.get("sw", [])):
                sheets["stars"].add(day(monday), cname, sid, name, house)
            if s.get("warn"):
                sheets["warnings"].add(cname, sid, name, house, WARNINGS[s["warn"]])
            if house in houses:
                houses[house] += points
            totals["passes"] += s.get("passes", 0)
            totals["month"] += sum(1 for t in log if t >= month)
            totals["fails"] += len(fails)
            totals["tests"] += len(s.get("tests", []))
            totals["merits"] += s.get("merits", 0)
            totals["warned"] += 1 if s.get("warn") else 0
            totals["points"] += points
            totals["q"] += marks
        try:
            first_day = day(state.get("fcd", ""))
        except ValueError:
            first_day = None
        sheets["classes"].add(cname, ", ".join(teachers.get(cid, [])), len(students), first_day, totals["passes"],
                              totals["month"], totals["fails"], totals["tests"],
                              round(sum(totals["q"]) / len(totals["q"]), 2) if totals["q"] else None,
                              totals["merits"], totals["warned"], totals["points"])
    for row in sorted(memo, key=lambda r: r[0], reverse=True):
        sheets["memo"].add(*row)
    for row in sorted(rewards, key=lambda r: r[0], reverse=True):
        sheets["rewards"].add(*row)
    sheets["quran"].rows.sort(key=lambda r: r[0], reverse=True)
    sheets["tests"].rows.sort(key=lambda r: r[0], reverse=True)
    sheets["stars"].rows.sort(key=lambda r: r[0], reverse=True)
    for u in users:
        status = "Turned off" if u["disabled"] else "Invited" if not u["password_hash"] else "Active"
        sheets["staff"].add(u["name"], u["email"], {"admin": "Administrator", "head": "HeadTeacher"}.get(u["role"], "Teacher"),
                            "All classes" if u["role"] in ("admin", "head") else
                            ", ".join(sorted((names.get(c, c) for c in member_of.get(u["id"], [])), key=natural)),
                            status, when(u["created_at"] * 1000).date() if u["created_at"] else None)
    figures = [("Classes", len(classes)), ("Students", len(sheets["students"].rows)),
               ("Staff accounts", len(users)), ("Memo passes", sum(1 for r in memo if r[4] == "Pass")),
               ("Quran scores", len(sheets["quran"].rows)), ("Level tests", len(sheets["tests"].rows)),
               *((f"{h} house points", n) for h, n in houses.items())]
    return list(sheets.values()), figures, now


def _fit(sheet, columns, rows):
    """Column widths from the header and the longest value, within sensible limits."""
    for i, (header, hint, fmt) in enumerate(columns):
        longest = max([len(header) + 3, hint] + [len(str(r[i])) + 1 for r in rows[:2000] if r[i] is not None and not fmt])
        sheet.column_dimensions[get_column_letter(i + 1)].width = min(longest, 48)


def plain_text(cells):
    """openpyxl saves any text that starts with "=" as a formula. Names and other details come from staff, so a
    student named "=HYPERLINK(...)" would become a live link or a formula in the admin's spreadsheet; they are kept
    as the text that was typed."""
    for cell in cells:
        if cell.data_type == "f" and isinstance(cell.value, str):
            cell.data_type, cell.quotePrefix = "s", True


def build(sheets, figures, made_at, made_by):
    book = Workbook()
    book.properties.creator = "MySafar"
    book.properties.title = "Safar Academy Enfield records"
    head_font, head_fill = Font(name="Calibri", bold=True, color=WHITE), PatternFill("solid", fgColor=NAVY)
    head_line = Border(bottom=Side(style="thin", color=NAVY))
    band = FormulaRule(formula=["MOD(ROW(),2)=1"], fill=PatternFill("solid", fgColor=BAND))

    cover = book.active
    cover.title = "Summary"
    cover.sheet_properties.tabColor = NAVY
    cover.sheet_view.showGridLines = False
    cover.column_dimensions["A"].width = 26
    cover.column_dimensions["B"].width = 72
    cover.page_setup.fitToWidth, cover.page_setup.fitToHeight = 1, 0
    cover.sheet_properties.pageSetUpPr.fitToPage = True
    cover["A1"] = "Safar Academy Enfield"
    cover["A1"].font = Font(name="Calibri", bold=True, size=18, color=NAVY)
    cover["A2"] = f"School records from MySafar, made {made_at.strftime('%-d %B %Y at %H:%M')} by {made_by}."
    cover["A2"].font = Font(name="Calibri", color=MUTED)
    cover["A4"] = "Key figures"
    cover["A4"].font = Font(name="Calibri", bold=True, size=13, color=NAVY)
    r = 5
    for label, value in figures:
        cover.cell(r, 1, label)
        cover.cell(r, 2, value).alignment = Alignment(horizontal="left")
        cover.cell(r, 2).number_format = "#,##0"
        r += 1
    r += 1
    cover.cell(r, 1, "Sheets").font = Font(name="Calibri", bold=True, size=13, color=NAVY)
    r += 1
    for s in sheets:
        link = cover.cell(r, 1, s.title)
        link.hyperlink = f"#'{s.title}'!A1"
        link.font = Font(name="Calibri", color=NAVY, underline="single")
        cover.cell(r, 2, f"{s.about} {len(s.rows):,} {'row' if len(s.rows) == 1 else 'rows'}.")
        r += 1
    r += 1
    cover.cell(r, 1, "Each sheet has filters on its top row: use the arrows to sort or pick what to show. "
                     "Dates and times are London time.").font = Font(name="Calibri", italic=True, color=MUTED)

    for s in sheets:
        ws = book.create_sheet(s.title)
        for i, (header, _, _) in enumerate(s.columns, 1):
            cell = ws.cell(1, i, header)
            cell.font, cell.fill, cell.border = head_font, head_fill, head_line
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.row_dimensions[1].height = 32
        formats = [fmt for _, _, fmt in s.columns]
        for row in s.rows:
            ws.append(row)
            plain_text(ws[ws.max_row])
        for i, fmt in enumerate(formats, 1):
            if fmt:
                for (cell,) in ws.iter_rows(min_row=2, min_col=i, max_col=i):
                    cell.number_format = fmt
        last = get_column_letter(len(s.columns))
        if s.rows:
            ws.auto_filter.ref = f"A1:{last}{len(s.rows) + 1}"
            ws.conditional_formatting.add(f"A2:{last}{len(s.rows) + 1}", band)
        else:
            ws.cell(2, 1, "Nothing recorded yet.").font = Font(name="Calibri", italic=True, color=MUTED)
        ws.freeze_panes = "A2"
        _fit(ws, s.columns, s.rows)
        ws.page_setup.orientation = "landscape"
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.print_title_rows = "1:1"
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


def sheet_list():
    """The sheets and what each holds, for the admin page."""
    return [("Summary", "Key figures and links to every sheet.")] + [(title, about) for title, about, _ in SHEETS.values()]


@bp.get("/admin/export")
@auth.admin_required
def export_page():
    return auth.page("admin_export.html", sheets=sheet_list())


@bp.get("/admin/export.xlsx")
@auth.admin_required
def export_download():
    with current_app.config["DB"].transaction() as db:
        sheets, figures, made_at = gather(db)
        auth.audit(db, "export.downloaded", None, {"classes": figures[0][1], "students": figures[1][1]})
    body = build(sheets, figures, made_at, g.user["name"])
    name = f"MySafar records {made_at.strftime('%Y-%m-%d %H%M')}.xlsx"
    return Response(body, mimetype=MIME, headers={"Content-Disposition": f'attachment; filename="{name}"'})
