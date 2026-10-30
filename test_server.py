"""Server tests.

The SQLite tests always run. The same tests run against PostgreSQL when
TEST_DATABASE_URL is set; each test uses its own temporary schema, which is
dropped afterwards, so existing tables in that database are never touched.
"""
import html
import io
import json
import os
import re
import runpy
import sqlite3
import tempfile
import time
import unittest
import uuid
from datetime import date, datetime
from contextlib import closing
from pathlib import Path
from unittest.mock import patch

import psycopg
from psycopg.conninfo import make_conninfo

import auth
import helpdesk
import server
import tv
from server import create_app, valid_state

ADMIN = ("admin@example.com", "bootstrap-password-1")


class FakeMailer:
    enabled = True
    sender_address = "dashboard@example.org"
    redirect_to = ""

    def __init__(self):
        self.sent = []
        self.html = []
        self.tags = []

    def send(self, to, subject, text, html=None, tag=None, background=False):
        self.sent.append((to, subject, text))
        self.html.append(html)
        self.tags.append(tag)
        return True

    def link(self):
        # The latest email with a link: a password-change notice may have followed it.
        text = next(t for _, _, t in reversed(self.sent) if re.search(r"/(?:invite|reset)/", t))
        return re.search(r"https?://\S+(/(?:invite|reset)/\S+)", text).group(1)


def csrf_of(response):
    return re.search(r'name="csrf" value="([^"]+)"', response.get_data(as_text=True)).group(1)


class DashboardTests:
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.state = {"S": [], "name": "Yr3-A", "fcd": "2026-09-13", "L": [9]}
        self.config = {"TESTING": True, "DATA_DIR": Path(self.temp.name) / "store",
                       "DATABASE_URL": self.database_url(), "COOKIE_SECURE": False,
                       "ADMIN_EMAIL": ADMIN[0], "DASHBOARD_PASSWORD": ADMIN[1]}
        self.app = create_app(self.config)
        # The first admin chooses their own password at first sign-in; tested separately.
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE users SET must_change = 0")
            # The historical house points the migration adds; test_historical_house_points checks them.
            self.seeded = [dict(r) for r in tx.execute("SELECT house, coin, n FROM house_points ORDER BY id").fetchall()]
            tx.execute("DELETE FROM house_points")
        self.client = self.sign_in(*ADMIN)
        self.headers = {"X-Dashboard-Request": "1"}

    def sign_in(self, email, password, app=None):
        client = (app or self.app).test_client()
        page = client.get("/login")
        data = {"csrf": csrf_of(page), "email": email, "password": password}
        client.last = client.post("/login", data=data)
        return client

    def add_user(self, email, role="teacher", password="teacher-password-1", classes=()):
        with self.app.config["DB"].transaction() as tx:
            user_id = tx.insert("INSERT INTO users (email, name, role, password_hash, created_at) VALUES (?, ?, ?, ?, ?)",
                                (email, email.split("@")[0].title(), role, auth.hasher.hash(password), 0))
            for class_id in classes:
                tx.execute("INSERT INTO classroom_members VALUES (?, ?)", (class_id, user_id))
        return user_id

    def query(self, sql, params=()):
        with self.app.config["DB"].transaction() as tx:
            return [dict(row) for row in tx.execute(sql, params).fetchall()]

    def records(self, client=None):
        return (client or self.client).get("/api/classes", headers=self.headers).get_json()

    def write(self, name="Yr3-A", revision=0, headers=None, state=None, class_id="yr3-a", client=None):
        state = dict(state or self.state, name=name)
        return (client or self.client).put("/api/classes/" + class_id, json={"json": json.dumps(state), "revision": revision},
                                           headers=self.headers if headers is None else headers)

    def delete(self, revision, class_id="yr3-a", client=None):
        return (client or self.client).delete("/api/classes/" + class_id, json={"revision": revision}, headers=self.headers)

    def form(self, client, page, action, **data):
        return client.post(action, data=dict(data, csrf=csrf_of(client.get(page))))

    def picture(self, size=(640, 480), fmt="PNG", **save):
        from PIL import Image
        out = io.BytesIO()
        Image.new("RGB", size, "#3366cc").save(out, fmt, **save)
        return out.getvalue()

    def upload_picture(self, client, data, json_reply=True):
        headers = {"Accept": "application/json"} if json_reply else {}
        return client.post("/account/avatar", headers=headers, content_type="multipart/form-data",
                           data={"csrf": csrf_of(client.get("/account")), "photo": (io.BytesIO(data), "me.png")})

    def test_profile_picture_default_is_first_initial(self):
        page = self.client.get("/account").get_data(as_text=True)
        self.assertIn('<span class="av" aria-hidden="true">A</span>', page)
        self.assertIn("Add a picture", page)
        self.assertNotIn("Remove picture", page)
        self.assertIn("script-src 'self'", self.client.get("/account").headers["Content-Security-Policy"])
        self.assertIn('"initial": "A", "avatar": null', self.client.get("/").get_data(as_text=True))

    def test_profile_picture_upload_is_reencoded_and_shown(self):
        from PIL import Image
        exif = Image.Exif()
        exif[0x010F] = "Secret Camera"
        reply = self.upload_picture(self.client, self.picture((900, 600), "JPEG", exif=exif.tobytes()))
        self.assertEqual(reply.status_code, 200, reply.data)
        url = reply.get_json()["url"]
        self.assertRegex(url, r"^/avatar/\d+\?v=\d+$")
        served = self.client.get(url)
        self.assertEqual(served.mimetype, "image/jpeg")
        self.assertIn("private", served.headers["Cache-Control"])
        with Image.open(io.BytesIO(served.data)) as image:
            self.assertEqual((image.format, image.size), ("JPEG", (256, 256)))
            self.assertNotIn(0x010F, image.getexif())
        self.assertNotIn(b"Secret Camera", served.data)
        # Shown by the name in the menu, on the dashboard, on Manage users and to other staff.
        self.assertIn(url, self.client.get("/account").get_data(as_text=True))
        self.assertIn(url, self.client.get("/").get_data(as_text=True))
        self.assertIn(url.replace("&", "&amp;"), self.client.get("/admin/users").get_data(as_text=True))
        self.add_user("teacher@example.com")
        teacher = self.sign_in("teacher@example.com", "teacher-password-1")
        self.assertEqual(teacher.get(url).status_code, 200)
        self.assertEqual(self.app.test_client().get(url).status_code, 302)
        self.assertEqual(len(self.query("SELECT 1 FROM audit_log WHERE action = 'avatar.changed'")), 1)
        # Removing it brings the initial back.
        self.form(self.client, "/account", "/account/avatar/remove")
        self.assertEqual(self.client.get(url).status_code, 404)
        self.assertIn('<span class="av" aria-hidden="true">A</span>', self.client.get("/account").get_data(as_text=True))

    def test_profile_picture_rejects_files_that_are_not_pictures(self):
        for data in (b"<svg xmlns='http://www.w3.org/2000/svg'/>", b"%PDF-1.4", b"", self.picture(fmt="BMP")):
            reply = self.upload_picture(self.client, data)
            self.assertEqual(reply.status_code, 400)
            self.assertFalse(reply.get_json()["ok"])
        self.assertEqual(self.query("SELECT COUNT(*) AS n FROM avatars")[0]["n"], 0)
        # A plain form post (no scripts) gets the Account page back with the error.
        reply = self.upload_picture(self.client, b"not a picture", json_reply=False)
        self.assertEqual(reply.status_code, 400)
        self.assertIn("couldn", reply.get_data(as_text=True))
        reply = self.upload_picture(self.client, self.picture(), json_reply=False)
        self.assertIn("Profile picture saved.", reply.get_data(as_text=True))
        self.assertIn("Remove picture", reply.get_data(as_text=True))

    def test_profile_picture_needs_form_token(self):
        reply = self.client.post("/account/avatar", content_type="multipart/form-data",
                                 data={"photo": (io.BytesIO(self.picture()), "me.png")})
        self.assertEqual(reply.status_code, 400)

    def test_first_admin_needs_admin_email(self):
        with self.assertRaisesRegex(RuntimeError, "ADMIN_EMAIL"):
            create_app(dict(self.config, DATABASE_URL=self.database_url(), DATA_DIR=Path(self.temp.name) / "new",
                            ADMIN_EMAIL=""))

    def test_starts_empty(self):
        self.assertEqual(self.records(), {})

    def test_health_public_but_dashboard_and_api_private(self):
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get("/healthz").get_json(), {"status": "ok"})
        for path in ("/", "/railway-storage.js", "/account", "/admin/users"):
            response = anonymous.get(path)
            self.assertEqual(response.status_code, 302)
            self.assertTrue(response.headers["Location"].startswith("/login"))
        self.assertEqual(anonymous.get("/api/classes").status_code, 401)
        self.assertEqual(anonymous.get("/login").status_code, 200)
        self.assertEqual(anonymous.get("/logo.png").mimetype, "image/png")
        response = self.client.get("/")
        self.assertEqual(response.status_code, 200)
        self.assertIn(b"window.RAILWAY=true", response.data)
        self.assertIn(b'"role": "admin"', response.data)
        self.assertIn(b"<main inert>", response.data)
        self.assertIn(b'/railway-storage.js', response.data)
        self.assertEqual(response.headers["Cache-Control"], "no-store")
        for path in ("/data/roster.csv", "/server.py", "/.env", "/Master_Dashboard.html", "/static/auth.css",
                     "/templates/login.html"):
            self.assertEqual(self.client.get(path).status_code, 404)

    def test_create_classroom_and_retry_is_idempotent(self):
        self.assertEqual(self.write().get_json(), {"revision": 1})
        self.assertEqual(self.write().get_json(), {"revision": 1})
        self.assertEqual(self.write("Someone else's Yr3-A").status_code, 409)
        self.assertEqual(json.loads(self.records()["yr3-a"]["json"])["name"], "Yr3-A")

    def test_save_survives_restart(self):
        self.write()
        self.assertEqual(self.write("Renamed", revision=1).status_code, 200)
        document = self.records(self.sign_in(*ADMIN, app=create_app(self.config)))["yr3-a"]
        self.assertEqual(json.loads(document["json"])["name"], "Renamed")
        self.assertEqual(document["revision"], 2)

    def test_conflicting_browser_cannot_overwrite_and_retry_is_idempotent(self):
        self.write()
        self.assertEqual(self.write("First edit", revision=1).status_code, 200)
        self.assertEqual(self.write("Stale edit", revision=1).status_code, 409)
        self.assertEqual(json.loads(self.records()["yr3-a"]["json"])["name"], "First edit")
        self.assertEqual(self.write("First edit", revision=1).status_code, 200)
        self.assertEqual(self.records()["yr3-a"]["revision"], 2)

    def test_delete_classroom(self):
        self.write()
        self.assertEqual(self.delete(0).status_code, 409)
        self.assertEqual(self.delete(1).status_code, 200)
        self.assertEqual(self.records(), {})
        self.assertEqual(self.delete(1).status_code, 200)
        self.assertEqual(self.write("Edited after delete", revision=1).status_code, 404)

    def test_quran_scores_from_zero_to_five(self):
        student = {"name": "A", "target": 1, "passes": 0, "merits": 0, "warn": 0,
                   "log": [], "lt": [], "lx": [], "ml": [], "q": {"2026-10-10": 0}}
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        student["q"] = {"2026-10-10": 6}
        self.assertEqual(self.write(revision=1, state=dict(self.state, S=[student])).status_code, 400)

    def test_quran_absent_is_minus_one(self):
        student = {"name": "A", "target": 1, "passes": 0, "merits": 0, "warn": 0,
                   "log": [], "lt": [], "lx": [], "ml": [], "q": {"2026-10-10": -1, "2026-10-11": 4}}
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        student["q"] = {"2026-10-10": -2}
        self.assertEqual(self.write(revision=1, state=dict(self.state, S=[student])).status_code, 400)

    def test_fails_and_surah_passes(self):
        student = {"name": "A", "target": 4, "passes": 2, "merits": 0, "warn": 0, "log": [1, 2], "lt": [3, 4],
                   "lx": ["", ""], "ml": [], "fl": [3], "q": {}}
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        book_done = dict(student, target=10, lt=[3, 10], lx=["", "Surah Al-Mulk"])
        self.assertEqual(self.write(revision=1, state=dict(self.state, S=[book_done])).status_code, 200)
        student["fl"] = ["x"]
        self.assertEqual(self.write(revision=2, state=dict(self.state, S=[student])).status_code, 400)

    def test_targets_run_on_through_the_juz_items(self):
        levels = self.state["L"]
        student = {"name": "A", "target": sum(levels) + 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [],
                   "lx": [], "ml": [], "q": {}}
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        student["target"] = sum(levels) + 136
        self.assertEqual(self.write(revision=1, state=dict(self.state, S=[student])).status_code, 400)
        self.assertEqual(self.write(revision=1, state=dict(self.state, Q=135, S=[student])).status_code, 200)
        student["target"] = sum(levels) + 137
        self.assertEqual(self.write(revision=2, state=dict(self.state, Q=135, S=[student])).status_code, 400)
        self.assertEqual(self.write(revision=2, state=dict(self.state, Q="135", S=[student])).status_code, 400)

    def test_student_details_from_csv(self):
        student = {"name": "A", "target": 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [],
                   "ml": [], "q": {}, "sid": "ST001", "gender": "Female", "year": "Year 3", "sy": "Year 4", "ist": "Yr3-A-AI-Sr Maryam"}
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        saved = json.loads(self.records()["yr3-a"]["json"])["S"][0]
        self.assertEqual((saved["sid"], saved["year"], saved["sy"]), ("ST001", "Year 3", "Year 4"))
        student["sid"] = 1
        self.assertEqual(self.write(revision=1, state=dict(self.state, S=[student])).status_code, 400)
        student["sid"], student["sy"] = "ST001", 4
        self.assertEqual(self.write(revision=1, state=dict(self.state, S=[student])).status_code, 400)

    def test_student_profile_details(self):
        student = {"name": "A", "target": 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [],
                   "ml": [], "q": {}, "sid": "SE012", "fid": "SP004", "dob": "2015-03-12", "ed": ["sid", "dob", "name", "class"],
                   "xid": ["SE011"], "xn": ["Aisha"]}
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        saved = json.loads(self.records()["yr3-a"]["json"])["S"][0]
        self.assertEqual((saved["fid"], saved["dob"], saved["ed"], saved["xid"]), ("SP004", "2015-03-12", ["sid", "dob", "name", "class"], ["SE011"]))
        for bad in ({"dob": "12/03/2015"}, {"dob": 2015}, {"fid": 4}, {"ed": ["q"]}, {"ed": "sid"}, {"xid": [1]}, {"xn": "A"}):
            self.assertEqual(self.write(revision=1, state=dict(self.state, S=[dict(student, **bad)])).status_code, 400)
        self.assertEqual(self.write(revision=1, state=dict(self.state, S=[dict(student, dob="")])).status_code, 200)

    def test_house_points(self):
        student = {"name": "A", "target": 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [],
                   "ml": [], "q": {}, "house": "Makkah", "tg": [1, 2], "tb": [3], "tr": [4]}
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        other = dict(student, name="B", house="Madinah", tg=[5], tb=[], tr=[])
        no_house = dict(student, name="C", house="", tg=[6] * 10)
        self.assertEqual(self.write("Yr4-B", class_id="yr4-b", state=dict(self.state, S=[other, no_house])).status_code, 200)
        self.assertEqual(self.client.get("/api/houses").get_json(), {"Makkah": 80, "Al-Aqsa": 0, "Madinah": 5})
        # Teachers see whole-school totals, but still only their own classrooms.
        self.add_user("t@example.com", classes=["yr4-b"])
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.assertEqual(teacher.get("/api/houses").get_json()["Makkah"], 80)
        self.assertEqual(list(self.records(teacher)), ["yr4-b"])
        self.assertEqual(self.app.test_client().get("/api/houses").status_code, 401)
        for bad in ({"house": "Hogwarts"}, {"house": None}, {"tg": "1"}, {"tb": [-1]}, {"tr": ["x"]}):
            self.assertEqual(self.write(revision=1, state=dict(self.state, S=[dict(student, **bad)])).status_code, 400)

    def test_last_student_number(self):
        student = {"name": "A", "target": 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [],
                   "ml": [], "q": {}}
        self.assertEqual(self.client.get("/api/last-sid").get_json(), {"last": 0})
        self.assertEqual(self.write(state=dict(self.state, S=[dict(student, sid="SE263"), dict(student, name="B", sid="ST900")])).status_code, 200)
        self.assertEqual(self.write("Yr4-B", class_id="yr4-b", state=dict(self.state, S=[dict(student, sid=" se99 ")])).status_code, 200)
        self.assertEqual(self.client.get("/api/last-sid").get_json(), {"last": 263})
        # Teachers get the whole school's number, though they only load their own classrooms.
        self.add_user("t@example.com", classes=["yr4-b"])
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").get("/api/last-sid").get_json(), {"last": 263})
        self.assertEqual(self.app.test_client().get("/api/last-sid").status_code, 401)

    def test_write_validation_and_cross_origin_protection(self):
        self.assertEqual(self.write(headers={}).status_code, 403)
        self.assertEqual(self.write(headers=dict(self.headers, **{"Sec-Fetch-Site": "cross-site"})).status_code, 403)
        self.assertEqual(self.write(headers=dict(self.headers, Origin="https://evil.example")).status_code, 403)
        self.assertEqual(self.write(client=self.app.test_client()).status_code, 401)
        for class_id in ("Yr3-A", "-yr3", "a" * 65, "yr3_a"):
            self.assertEqual(self.write(class_id=class_id).status_code, 400)
        for document in ({}, {"json": "not json", "revision": 0}, {"json": "{}", "revision": 0},
                         {"json": json.dumps(self.state), "revision": -1},
                         {"json": json.dumps(dict(self.state, S=[{"name": "Invalid student"}])), "revision": 0}):
            response = self.client.put("/api/classes/yr3-a", json=document, headers=self.headers)
            self.assertEqual(response.status_code, 400)
        self.assertEqual(self.client.delete("/api/classes/yr3-a", headers=self.headers).status_code, 403)
        self.assertEqual(self.records(), {})

    def test_qaidah_tests_and_stars_are_checked(self):
        student = {"name": "Aisha", "target": 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [],
                   "ml": [], "fl": [], "q": {}}
        state = dict(self.state, L=[9, 9, 9, 9, 10, 10, 10, 12, 15, 14])
        self.assertTrue(valid_state(dict(state, S=[student])))
        self.assertTrue(valid_state(dict(state, S=[dict(student, qa=20)])))
        for bad in (-1, 21, "3", 2.5):
            self.assertFalse(valid_state(dict(state, S=[dict(student, qa=bad)])))
        test = {"d": "2026-10-14", "k": "m", "n": 12, "sc": 4, "ok": True}
        self.assertTrue(valid_state(dict(state, S=[dict(student, tests=[test, dict(test, k="q", n=3, ok=False, pq=2)])])))
        # A pass keeps the time its blue token was given, so taking the test back takes the token back too.
        self.assertTrue(valid_state(dict(state, S=[dict(student, tb=[1760400000000], tests=[dict(test, tk=1760400000000)])])))
        # A surah chosen by item number on the Record page is kept as written, with no token.
        self.assertTrue(valid_state(dict(state, S=[dict(student, tests=[dict(test, x="212. Surah Al-Mulk")])])))
        for bad in ({"d": "14/10/2026"}, {"k": "x"}, {"n": 0}, {"sc": 6}, {"ok": 1}, {"pq": 21}, {"tk": -1}, {"tk": "now"},
                    {"x": 212}, {"x": "x" * 81}):
            self.assertFalse(valid_state(dict(state, S=[dict(student, tests=[dict(test, **bad)])])))
        self.assertFalse(valid_state(dict(state, S=[dict(student, tests={})])))
        # House points from the Record page keep what they were for, with the time their token was given.
        self.assertTrue(valid_state(dict(state, S=[dict(student, tr=[1760400000000], hp=[{"t": 1760400000000, "r": "bk"}])])))
        for bad in ({}, [{"t": 1760400000000}], [{"t": -1, "r": "bk"}], [{"t": "now", "r": "bk"}], [{"t": 1, "r": "xx"}]):
            self.assertFalse(valid_state(dict(state, S=[dict(student, hp=bad)])))
        # Book passes: when and which book, the Qaidah Book or the Tajweed Book.
        self.assertTrue(valid_state(dict(state, S=[dict(student, bp=[{"t": 1760400000000, "b": "q"}, {"t": 1760400000001, "b": "t"}])])))
        for bad in ({}, [{"t": 1760400000000}], [{"t": -1, "b": "q"}], [{"t": "now", "b": "t"}], [{"t": 1, "b": "x"}]):
            self.assertFalse(valid_state(dict(state, S=[dict(student, bp=bad)])))
        self.assertTrue(valid_state(dict(state, S=[dict(student, sw=["2026-10-12"])])))
        self.assertFalse(valid_state(dict(state, S=[dict(student, sw=["12 Oct"])])))

    def test_classroom_limit(self):
        with patch("server.MAX_CLASSES", 2):
            self.assertEqual(self.write(class_id="a").status_code, 200)
            self.assertEqual(self.write(class_id="b").status_code, 200)
            self.assertEqual(self.write(class_id="c").status_code, 400)
            self.assertEqual(self.write(revision=1, class_id="a").status_code, 200)

    def test_earlier_classes_table_is_deleted(self):
        config = dict(self.config, DATABASE_URL=self.database_url(), DATA_DIR=Path(self.temp.name) / "earlier")
        self.execute(config, "CREATE TABLE classes (id TEXT PRIMARY KEY, json TEXT NOT NULL, revision INTEGER NOT NULL)",
                     "INSERT INTO classes VALUES ('y1a', '{}', 4)")
        create_app(config)
        self.assertFalse(self.table_exists(config, "classes"))
        create_app(config)  # later startups have nothing to delete

    def test_existing_records_are_kept_when_accounts_are_added(self):
        config = dict(self.config, DATABASE_URL=self.database_url(), DATA_DIR=Path(self.temp.name) / "upgrade")
        # The database as the shared-login version left it.
        self.execute(config, "CREATE TABLE classrooms (id TEXT PRIMARY KEY, json TEXT NOT NULL, revision INTEGER NOT NULL)",
                     "INSERT INTO classrooms VALUES ('yr3-a', '{\"name\": \"Yr3-A\"}', 7)")
        app = create_app(config)
        with app.config["DB"].transaction() as tx:
            tx.execute("UPDATE users SET must_change = 0")
        self.assertEqual(self.records(self.sign_in(*ADMIN, app=app))["yr3-a"]["revision"], 7)

    # --- Signing in ---

    def test_first_admin_must_choose_a_new_password(self):
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE users SET must_change = 1")
        client = self.sign_in(*ADMIN)
        self.assertEqual(client.get("/").headers["Location"], "/account")
        self.assertEqual(client.get("/api/classes").status_code, 403)
        weak = self.form(client, "/account", "/account/password", current=ADMIN[1], password="short", confirm="short")
        self.assertEqual(weak.status_code, 400)
        good = self.form(client, "/account", "/account/password", current=ADMIN[1],
                         password="a brand new passphrase", confirm="a brand new passphrase")
        self.assertEqual(good.status_code, 200)
        self.assertEqual(client.get("/").status_code, 200)
        self.assertEqual(self.sign_in(*ADMIN).last.status_code, 401)
        self.assertEqual(self.sign_in(ADMIN[0], "a brand new passphrase").last.status_code, 302)

    def test_wrong_password_unknown_email_and_lockout(self):
        self.add_user("t@example.com")
        wrong = self.sign_in("t@example.com", "not the password")
        unknown = self.sign_in("nobody@example.com", "not the password")
        self.assertEqual((wrong.last.status_code, unknown.last.status_code), (401, 401))
        self.assertIn(b"don&#39;t match", wrong.last.data)
        self.assertIn(b"don&#39;t match", unknown.last.data)
        for _ in range(4):
            self.sign_in("t@example.com", "still wrong")
        # Locked out for 15 minutes, even with the right password.
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").last.status_code, 429)
        self.assertEqual(self.sign_in(*ADMIN).last.status_code, 302)
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE attempts SET at = at - 1000")
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").last.status_code, 302)

    def test_sign_in_needs_the_form_token(self):
        client = self.app.test_client()
        response = client.post("/login", data={"email": ADMIN[0], "password": ADMIN[1]})
        self.assertEqual(response.status_code, 400)

    def test_sign_in_returns_to_the_page_asked_for_but_never_another_site(self):
        self.assertEqual(self.app.test_client().get("/admin/users").headers["Location"], "/login?next=/admin/users")
        for target, expected in (("/admin/users", "/admin/users"), ("//evil.example", "/"),
                                 ("https://evil.example", "/"), ("/\\evil.example", "/")):
            client = self.app.test_client()
            page = client.get("/login")
            response = client.post("/login", data={"csrf": csrf_of(page), "email": ADMIN[0], "password": ADMIN[1],
                                                   "next": target})
            self.assertEqual(response.headers["Location"], expected)

    def test_new_tab_goes_to_sign_in(self):
        # The first dashboard load after signing in is marked fresh (greeting, tab marked as signed in); later
        # loads are not, so a tab without the mark signs out in the browser (TAB_GUARD).
        client = self.sign_in(*ADMIN)
        first = client.get("/").get_data(as_text=True)
        self.assertIn('"fresh": true', first)
        self.assertIn("mysafar.tab", first)
        self.assertIn('"fresh": false', client.get("/").get_data(as_text=True))

    def test_sign_out(self):
        response = self.form(self.client, "/account", "/logout")
        self.assertEqual(response.headers["Location"], "/login?out=1")
        self.assertEqual(self.client.get("/api/classes").status_code, 401)

    def test_sessions_expire(self):
        client = self.sign_in(*ADMIN)
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE sessions SET last_seen_at = last_seen_at - ?", (auth.SESSION_IDLE + 1,))
        self.assertEqual(client.get("/api/classes").status_code, 401)
        # No "keep me signed in": the cookie ends with the browser, and the sign-in after 20 idle minutes.
        self.assertEqual(auth.SESSION_IDLE, 20 * 60)
        self.assertNotIn('name="remember"', self.app.test_client().get("/login").get_data(as_text=True))
        client = self.sign_in(*ADMIN)
        self.assertNotIn("Max-Age", client.last.headers["Set-Cookie"])
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE sessions SET expires_at = ?", (int(time.time()) - 1,))
        self.assertEqual(client.get("/api/classes").status_code, 401)
        # A 30-day sign-in left from the old option ends at once.
        client = self.sign_in(*ADMIN)
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE sessions SET expires_at = ?", (int(time.time()) + 30 * 86400,))
        self.assertEqual(client.get("/api/classes").status_code, 401)

    def test_behind_railways_proxy(self):
        app = create_app(dict(self.config, COOKIE_SECURE=True, TRUST_PROXY=True,
                              PUBLIC_URL="https://dash.example.org"))
        client = app.test_client()
        railway = {"X-Forwarded-Proto": "https", "X-Forwarded-Host": "dash.example.org", "X-Real-IP": "203.0.113.9"}
        page = client.get("/login", headers=railway)
        # The browser sends the public site as Origin; the app sees an internal Host.
        response = client.post("/login", headers=dict(railway, Origin="https://dash.example.org"),
                               data={"csrf": csrf_of(page), "email": ADMIN[0], "password": ADMIN[1]})
        self.assertEqual(response.status_code, 302)
        self.assertIn({"ip": "203.0.113.9"}, self.query("SELECT ip FROM sessions"))
        evil = client.put("/api/classes/yr3-a", json={}, headers=dict(railway, Origin="https://evil.example",
                                                                       **self.headers))
        self.assertEqual(evil.status_code, 403)

    def test_secure_cookies(self):
        app = create_app(dict(self.config, COOKIE_SECURE=True))
        client = app.test_client()
        page = client.get("/login", base_url="https://localhost")
        response = client.post("/login", base_url="https://localhost",
                               data={"csrf": csrf_of(page), "email": ADMIN[0], "password": ADMIN[1]})
        cookie = response.headers["Set-Cookie"]
        for part in ("__Host-safar=", "Secure", "HttpOnly", "SameSite=Lax", "Path=/"):
            self.assertIn(part, cookie)
        self.assertIn("max-age=31536000", response.headers["Strict-Transport-Security"])

    # --- Teachers and admins ---

    def test_teachers_see_and_edit_only_their_classrooms(self):
        self.write(class_id="yr3-a")
        self.write("Yr4-B", class_id="yr4-b")
        self.add_user("t@example.com", classes=["yr3-a"])
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.assertEqual(list(self.records(teacher)), ["yr3-a"])
        self.assertIn(b'"role": "teacher"', teacher.get("/").data)
        self.assertEqual(self.write("Mine", revision=1, client=teacher).status_code, 200)
        self.assertEqual(self.write("Not mine", revision=1, class_id="yr4-b", client=teacher).status_code, 403)
        self.assertEqual(self.write("New", class_id="yr5-c", client=teacher).status_code, 403)
        self.assertEqual(self.delete(2, client=teacher).status_code, 403)
        for path in ("/admin/users", "/admin/history", "/admin/activity"):
            self.assertEqual(teacher.get(path).status_code, 403)
        self.assertEqual(json.loads(self.records()["yr4-b"]["json"])["name"], "Yr4-B")
        # The Students page is for administrators: the account pages' menu links to it only for them.
        self.assertNotIn(b'href="/#students"', teacher.get("/account").data)
        self.assertIn(b'href="/#students"', self.client.get("/account").data)
        # So is Rankings (classes by Memo passes).
        self.assertNotIn(b'href="/#rankings"', teacher.get("/account").data)
        self.assertIn(b'href="/#rankings"', self.client.get("/account").data)

    def test_tutorial_starts_once_for_everyone(self):
        self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.assertIn('"tour": true', teacher.get("/").get_data(as_text=True))
        # Admins get it too, and everyone can retake it from Account.
        self.assertIn('"tour": true', self.client.get("/").get_data(as_text=True))
        self.assertIn(b'href="/#tour"', teacher.get("/account").data)
        self.assertIn(b'href="/#tour"', self.client.get("/account").data)
        # Finishing or skipping it is a dashboard request like any other write.
        self.assertEqual(teacher.post("/api/tour", json={}).status_code, 403)
        self.assertEqual(teacher.post("/api/tour", json={}, headers=self.headers).status_code, 200)
        self.assertIn('"tour": false', teacher.get("/").get_data(as_text=True))
        self.assertTrue(self.query("SELECT tour_done FROM users WHERE email = ?", ("t@example.com",))[0]["tour_done"])

    def test_admin_downloads_the_school_spreadsheet(self):
        from io import BytesIO
        from openpyxl import load_workbook
        t = int(datetime(2026, 10, 14, 17, 30).timestamp() * 1000)
        student = {"name": "Musa Khan", "sid": "ST099", "fid": "SP040", "dob": "2016-05-02", "house": "Madinah", "target": 12, "passes": 2, "merits": 1,
                   "warn": 2, "log": [t - 864e5, t], "lt": [10, 11], "lx": ["", ""], "ml": [t], "fl": [t], "qa": 3,
                   "q": {"2026-10-12": -1, "2026-10-13": 4, "2026-10-14": 5}, "tg": [t, t + 1], "tb": [t], "tr": [], "sw": ["2026-10-12"],
                   "hp": [{"t": t + 1, "r": "hw"}],
                   "tests": [{"d": "2026-10-14", "k": "m", "n": 11, "sc": 4, "ok": True, "tk": t}]}
        self.write(state=dict(self.state, S=[student], L=[9, 9]), class_id="yr3-a")
        self.add_user("t@example.com", classes=["yr3-a"])
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        for path in ("/admin/export", "/admin/export.xlsx"):
            self.assertEqual(teacher.get(path).status_code, 403)
        self.assertIn(b'href="/admin/export.xlsx"', self.client.get("/admin/export").data)
        response = self.client.get("/admin/export.xlsx")
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response.mimetype, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
        self.assertIn("attachment", response.headers["Content-Disposition"])
        book = load_workbook(BytesIO(response.data))
        self.assertEqual(book.sheetnames, ["Summary", "Students", "Classes", "Quran scores", "Memo passes and fails",
                                           "Level tests", "Rewards", "Stars of the week", "Warnings", "Staff"])
        rows = {name: [[c.value for c in r] for r in book[name].iter_rows()] for name in book.sheetnames}
        students = dict(zip(rows["Students"][0], rows["Students"][1]))
        self.assertEqual((students["Student"], students["Memo target"], students["Memo passes"], students["Warning"],
                          students["House points"], students["Quran average"]), ("Musa Khan", "L2 T3", 2, "Warning 2", 30, 4.5))
        self.assertEqual(rows["Classes"][1][:3], ["Yr3-A", "T", 1])
        self.assertEqual([r[4] for r in rows["Quran scores"][1:]], [5, 4, "Absent"])
        self.assertEqual([(r[4], r[5]) for r in rows["Memo passes and fails"][1:]],
                         [("Pass", "L2 T2"), ("Not passed", None), ("Pass", "L2 T1")])
        self.assertEqual(rows["Level tests"][1][5:], ["L2 T2", 4, "Passed", "Yes"])
        self.assertEqual(sorted(r[5] for r in rows["Rewards"][1:]),
                         ["Blue token", "Green token", "Green token: Additional class or homework", "Merit"])
        self.assertEqual(book["Students"].freeze_panes, "A2")
        self.assertEqual(book["Students"].auto_filter.ref, "A1:Y2")
        self.assertEqual((students["Family ID"], students["Date of birth"].date()), ("SP040", date(2016, 5, 2)))
        self.assertIn(["Admin@Example.com".lower(), "Administrator"], [r[1:3] for r in rows["Staff"][1:]])
        self.assertEqual(self.query("SELECT action FROM audit_log WHERE action = 'export.downloaded'"),
                         [{"action": "export.downloaded"}])

    def test_spreadsheet_names_targets_past_the_book(self):
        import export
        self.assertEqual(export.target_label(1, [9]), "L1 T1")
        self.assertEqual(export.target_label(10, [9]), "198 An-Naba")
        self.assertEqual(export.target_label(10 + len(export.JUZ_ITEMS), [9]), "Quran complete")

    def test_only_admins_add_or_remove_students(self):
        student = {"name": "Musa Khan", "target": 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [],
                   "ml": [], "fl": [], "q": {}, "tg": [], "tb": [], "tr": [], "house": "Madinah", "sid": "ST099"}
        self.write(class_id="yr3-a", state=dict(self.state, S=[student]))
        self.add_user("t@example.com", classes=["yr3-a"])
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        other = dict(student, name="Aisha Ali", sid="ST100")
        for students in ([student, other], []):
            response = self.write(state=dict(self.state, S=students), revision=1, client=teacher)
            self.assertEqual(response.status_code, 403)
        self.assertEqual(len(json.loads(self.records(teacher)["yr3-a"]["json"])["S"]), 1)
        # Teachers still record work for the students they have.
        self.assertEqual(self.write(state=dict(self.state, S=[dict(student, merits=1)]), revision=1, client=teacher).status_code, 200)
        self.assertEqual(self.write(state=dict(self.state, S=[dict(student, merits=1), other]), revision=2).status_code, 200)
        # House points are given by admins only: a teacher's save can't add a coin, an admin's can.
        response = self.write(state=dict(self.state, S=[dict(student, merits=1, tg=[1760400000000]), other]), revision=3, client=teacher)
        self.assertEqual(response.status_code, 403)
        self.assertEqual(self.write(state=dict(self.state, S=[dict(student, merits=1, tg=[1760400000000]), other]), revision=3).status_code, 200)
        self.assertEqual(self.write(state=dict(self.state, S=[dict(student, merits=2, tg=[1760400000000]), other]), revision=4, client=teacher).status_code, 200)

    def test_invite_without_email(self):
        self.write()
        response = self.form(self.client, "/admin/users", "/admin/users/invite", name="Sr Maryam",
                             email="Maryam@Example.com", role="teacher", classes=["yr3-a", "not-a-class"])
        self.assertEqual(response.status_code, 200)
        link = re.search(r'value="https?://[^/"]+(/invite/[^"]+)"', response.get_data(as_text=True)).group(1)
        new = self.app.test_client()
        self.assertIn(b"Welcome, Sr Maryam", new.get(link).data)
        page = new.get(link)
        self.assertEqual(new.post(link, data={"csrf": csrf_of(page), "password": "short", "confirm": "short"}).status_code, 400)
        page = new.get(link)
        response = new.post(link, data={"csrf": csrf_of(page), "password": "olive tree morning", "confirm": "olive tree morning"})
        self.assertEqual(response.headers["Location"], "/")
        self.assertEqual(list(self.records(new)), ["yr3-a"])
        self.assertEqual(self.app.test_client().get(link).status_code, 410)
        self.assertEqual(self.sign_in("maryam@example.com", "olive tree morning").last.status_code, 302)
        duplicate = self.form(self.client, "/admin/users", "/admin/users/invite", name="Again",
                              email="maryam@example.com", role="teacher")
        self.assertEqual(duplicate.status_code, 400)

    def test_invite_and_reset_by_email(self):
        mailer = FakeMailer()
        self.app.config["MAILER"] = mailer
        response = self.form(self.client, "/admin/users", "/admin/users/invite", name="Br Ahmed",
                             email="ahmed@example.com", role="teacher")
        self.assertNotIn(b"/invite/", response.data)
        self.assertEqual(mailer.sent[-1][0], "ahmed@example.com")
        self.assertIn("72 hours", mailer.sent[-1][2])
        new = self.app.test_client()
        page = new.get(mailer.link())
        new.post(mailer.link(), data={"csrf": csrf_of(page), "password": "olive tree morning", "confirm": "olive tree morning"})
        # Forgotten password: the same answer whether or not there is an account.
        anonymous = self.app.test_client()
        unknown = self.form(anonymous, "/forgot", "/forgot", email="nobody@example.com")
        self.assertEqual(len(mailer.sent), 1)
        known = self.form(anonymous, "/forgot", "/forgot", email="Ahmed@example.com")
        self.assertEqual(unknown.data, known.data)
        self.assertEqual(len(mailer.sent), 2)
        self.assertIn("1 hour", mailer.sent[-1][2])
        page = anonymous.get(mailer.link())
        anonymous.post(mailer.link(), data={"csrf": csrf_of(page), "password": "river stone evening",
                                            "confirm": "river stone evening"})
        # Resetting signs out every other session.
        self.assertEqual(new.get("/api/classes").status_code, 401)
        self.assertEqual(self.sign_in("ahmed@example.com", "river stone evening").last.status_code, 302)
        # The reset is followed by a notice that the password changed.
        self.assertEqual((len(mailer.sent), mailer.tags[-1]), (3, "security"))
        self.assertIn("reset with a reset link", mailer.sent[-1][2])
        for _ in range(3):
            self.form(anonymous, "/forgot", "/forgot", email="ahmed@example.com")
        self.assertEqual(len(mailer.sent), 5)  # limited to 3 a quarter of an hour

    def test_emails_have_html_and_text_parts(self):
        self.write()
        mailer = FakeMailer()
        self.app.config.update(MAILER=mailer, PUBLIC_URL="https://mysafar.example", SUPPORT_EMAIL="office@example.com")
        self.form(self.client, "/admin/users", "/admin/users/invite", name="Sr <b>Maryam</b>",
                  email="maryam@example.com", role="teacher", classes=["yr3-a"])
        to, subject, text = mailer.sent[-1]
        html = mailer.html[-1]
        self.assertEqual((to, mailer.tags[-1]), ("maryam@example.com", "invite"))
        self.assertIn("invited", subject)
        link = "https://mysafar.example" + mailer.link()
        for part in (text, html):
            self.assertIn(link, part)
            self.assertIn("Yr3-A", part)  # the classrooms they were given
            self.assertIn("Teacher", part)  # their role
            self.assertIn("office@example.com", part)
        self.assertIn("Sr <b>Maryam</b>", text)
        # Names are escaped in the HTML part.
        self.assertIn("Sr &lt;b&gt;Maryam&lt;/b&gt;", html)
        self.assertNotIn("<b>Maryam", html)
        self.assertIn('src="https://mysafar.example/logo.png"', html)
        self.assertIn("72 hours", html)
        # The admin who sent it is named.
        self.assertIn("Administrator has given you an account", text)

    def test_changing_a_password_sends_a_notice(self):
        mailer = FakeMailer()
        self.app.config["MAILER"] = mailer
        self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.form(teacher, "/account", "/account/password", current="wrong-password-1", password="olive tree morning",
                  confirm="olive tree morning")
        self.assertEqual(mailer.sent, [])
        self.form(teacher, "/account", "/account/password", current="teacher-password-1",
                  password="olive tree morning", confirm="olive tree morning")
        to, subject, text = mailer.sent[-1]
        self.assertEqual((to, mailer.tags[-1]), ("t@example.com", "security"))
        self.assertIn("password was changed", subject)
        self.assertIn("changed on the Account page", text)
        self.assertIn("/forgot", text)

    def test_admin_sends_a_test_email_and_previews_the_emails(self):
        mailer = FakeMailer()
        self.app.config["MAILER"] = mailer
        page = self.client.get("/admin/users").get_data(as_text=True)
        self.assertIn("dashboard@example.org", page)
        response = self.form(self.client, "/admin/users", "/admin/email/test")
        self.assertIn(b"Test email sent to admin@example.com", response.data)
        self.assertEqual((mailer.sent[-1][0], mailer.tags[-1]), (ADMIN[0], "test"))
        self.assertIn("email.test", [r["action"] for r in self.query("SELECT action FROM audit_log")])
        for name in ("invite", "reset", "password_changed", "test"):
            preview = self.client.get(f"/admin/email/preview/{name}")
            self.assertEqual(preview.status_code, 200)
            self.assertIn("style-src 'unsafe-inline'", preview.headers["Content-Security-Policy"])
            self.assertNotIn("script", preview.headers["Content-Security-Policy"])
            self.assertIn(b"Safar Academy", preview.data)
            text = self.client.get(f"/admin/email/preview/{name}?text=1")
            self.assertEqual(text.mimetype, "text/plain")
            self.assertTrue(text.get_data(as_text=True).startswith("Subject: "))
        self.assertEqual(self.client.get("/admin/email/preview/other").status_code, 404)
        self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.assertEqual(teacher.get("/admin/email/preview/invite").status_code, 403)
        self.assertEqual(self.form(teacher, "/account", "/admin/email/test").status_code, 403)
        self.assertNotIn("unsafe-inline", self.client.get("/admin/users").headers["Content-Security-Policy"])

    def test_test_email_needs_email_set_up(self):
        response = self.form(self.client, "/admin/users", "/admin/email/test")
        self.assertEqual(response.status_code, 400)
        self.assertIn(b"isn&#39;t set up", self.client.get("/admin/users").data.replace(b"'", b"&#39;"))

    def test_admin_resets_a_teachers_password_without_seeing_it(self):
        user_id = self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        stored = self.query("SELECT password_hash FROM users WHERE id = ?", (user_id,))[0]["password_hash"]
        self.assertEqual(self.form(teacher, "/account", f"/admin/users/{user_id}/link").status_code, 403)
        response = self.form(self.client, "/admin/users", f"/admin/users/{user_id}/link")
        text = response.get_data(as_text=True)
        link = re.search(r'value="https?://[^/"]+(/reset/[^"]+)"', text).group(1)
        for page in (text, self.client.get("/admin/users").get_data(as_text=True),
                     self.client.get("/admin/activity").get_data(as_text=True)):
            self.assertNotIn(stored, page)
            self.assertNotIn("teacher-password-1", page)
        # Nothing changes until the teacher chooses a new password.
        self.assertEqual(teacher.get("/api/classes").status_code, 200)
        new = self.app.test_client()
        page = new.get(link)
        new.post(link, data={"csrf": csrf_of(page), "password": "river stone evening", "confirm": "river stone evening"})
        self.assertEqual(teacher.get("/api/classes").status_code, 401)
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").last.status_code, 401)
        self.assertEqual(self.sign_in("t@example.com", "river stone evening").last.status_code, 302)
        self.assertEqual(self.app.test_client().get(link).status_code, 410)
        self.assertEqual(self.client.get("/api/classes", headers=self.headers).status_code, 200)

    def test_reset_links_expire(self):
        self.add_user("t@example.com")
        mailer = FakeMailer()
        self.app.config["MAILER"] = mailer
        self.form(self.app.test_client(), "/forgot", "/forgot", email="t@example.com")
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE user_tokens SET expires_at = ?", (int(time.time()) - 1,))
        self.assertEqual(self.app.test_client().get(mailer.link()).status_code, 410)

    def test_turning_off_an_account_signs_it_out(self):
        user_id = self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.form(self.client, "/admin/users", f"/admin/users/{user_id}/access", disable="1")
        self.assertEqual(teacher.get("/api/classes").status_code, 401)
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").last.status_code, 401)
        self.form(self.client, "/admin/users", f"/admin/users/{user_id}/access", disable="0")
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").last.status_code, 302)

    def test_the_last_admin_stays_an_admin(self):
        admin_id = self.query("SELECT id FROM users")[0]["id"]
        response = self.form(self.client, "/admin/users", f"/admin/users/{admin_id}/update", name="Me", role="teacher")
        self.assertEqual(response.status_code, 400)
        response = self.form(self.client, "/admin/users", f"/admin/users/{admin_id}/access", disable="1")
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.query("SELECT role, disabled FROM users"), [{"role": "admin", "disabled": 0}])

    def test_change_classrooms_and_role(self):
        self.write()
        user_id = self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.assertEqual(self.records(teacher), {})
        self.form(self.client, "/admin/users", f"/admin/users/{user_id}/update", name="T", role="teacher", classes=["yr3-a"])
        self.assertEqual(list(self.records(teacher)), ["yr3-a"])
        self.form(self.client, "/admin/users", f"/admin/users/{user_id}/update", name="T", role="admin")
        self.assertEqual(self.delete(1, client=teacher).status_code, 200)

    # --- History and activity ---

    def test_history_restores_an_earlier_version_and_a_deleted_classroom(self):
        self.write()
        self.write("Renamed", revision=1)
        # Saves by the same person within 15 minutes share one history entry.
        self.assertEqual(len(self.query("SELECT id FROM classroom_history")), 1)
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE classroom_history SET saved_at = saved_at - 3600, json = ?", (json.dumps(self.state),))
        self.write("Renamed again", revision=2)
        self.assertEqual(self.delete(3).status_code, 200)
        self.assertEqual(self.records(), {})
        listing = self.client.get("/admin/history").get_data(as_text=True)
        self.assertIn("Deleted", listing)
        page = self.client.get("/admin/history/yr3-a")
        first = self.query("SELECT id FROM classroom_history ORDER BY id")[0]["id"]
        response = self.client.post("/admin/history/yr3-a/restore", data={"csrf": csrf_of(page), "entry": first})
        self.assertEqual(response.status_code, 200)
        record = self.records()["yr3-a"]
        self.assertEqual((json.loads(record["json"])["name"], record["revision"]), ("Yr3-A", 1))

    def test_clear_removes_every_progress_record(self):
        t = 1760000000000
        tests = [{"d": "2026-10-13", "k": "m", "n": 3, "sc": 4, "ok": True, "tk": t},
                 {"d": "2026-10-13", "k": "m", "n": 4, "sc": 1, "ok": False},
                 {"d": "2026-10-13", "k": "q", "n": 2, "sc": 5, "ok": True, "pq": 1, "tk": t}]
        amina = self.tv_student("Amina", target=4, passes=3, merits=2, warn=1, log=[t, t + 1], lt=[2, 3], lx=["", ""],
                                ml=[t, t], fl=[t], q={"2026-10-13": 4}, tb=[t, t], qa=2, sw=["2026-10-12"], tests=tests)
        self.write(state=dict(self.state, S=[amina, self.tv_student("Bilal")]))
        page = self.client.get("/admin/clear-test-data").get_data(as_text=True)
        self.assertIn("Yr3-A", page)
        self.assertEqual(self.client.get("/admin/clear-test-data/backup.json").get_json()["classrooms"]["yr3-a"]["revision"], 1)
        self.assertIn("Type CLEAR", self.form(self.client, "/admin/clear-test-data", "/admin/clear-test-data",
                                              confirm="no").get_data(as_text=True))
        self.assertEqual(self.records()["yr3-a"]["revision"], 1)
        done = self.form(self.client, "/admin/clear-test-data", "/admin/clear-test-data", confirm="clear")
        self.assertIn("Cleared 1 class", done.get_data(as_text=True))
        record = self.records()["yr3-a"]
        state = json.loads(record["json"])
        cleared = state["S"][0]
        self.assertEqual(record["revision"], 2)
        self.assertTrue(valid_state(state))
        for key, empty in (("passes", 0), ("log", []), ("lt", []), ("lx", []), ("sw", []), ("tests", []), ("merits", 0),
                           ("ml", []), ("warn", 0), ("fl", []), ("q", {}), ("tb", []), ("tg", []), ("tr", [])):
            self.assertEqual(cleared[key], empty, key)
        self.assertEqual((cleared["name"], cleared["target"], cleared["qa"]), ("Amina", 1, 0))
        self.assertEqual(self.client.get("/api/houses").get_json(), {"Makkah": 0, "Al-Aqsa": 0, "Madinah": 0})
        self.assertEqual(json.loads(self.query("SELECT json FROM classroom_history ORDER BY id")[-2]["json"])["S"][0]["passes"], 3)
        self.assertIn("no progress records", self.client.get("/admin/clear-test-data").get_data(as_text=True))
        self.add_user("t@example.com")
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").get("/admin/clear-test-data").status_code, 403)

    def test_image_copies_every_module(self):
        # The Dockerfile names each file it copies; a module left out stops the live site from starting.
        copied = set(re.search(r"^COPY (server\.py .*) \./$", Path("Dockerfile").read_text(), re.M).group(1).split())
        local = {p.name for p in Path(".").glob("*.py") if not p.name.startswith("test_")}
        imported = {m + ".py" for m in re.findall(r"^import (\w+)$", Path("server.py").read_text() + Path("tv.py").read_text(), re.M)}
        self.assertEqual((imported & local) - copied, set())

    def test_historical_house_points(self):
        points = {"tg": 5, "tb": 20, "tr": 50}
        seeded = {}
        for r in self.seeded:
            seeded[r["house"]] = seeded.get(r["house"], 0) + r["n"] * points[r["coin"]]
        self.assertEqual(seeded, {"Al-Aqsa": 130, "Madinah": 220, "Makkah": 240})
        self.write(state=dict(self.state, S=[self.tv_student("Amina", house="Makkah", tg=[1, 2])]))
        self.assertIn("None added yet", self.client.get("/admin/house-points").get_data(as_text=True))
        bad = self.form(self.client, "/admin/house-points", "/admin/house-points", house="Makkah", tg="x")
        self.assertEqual(bad.status_code, 400)
        done = self.form(self.client, "/admin/house-points", "/admin/house-points", house="Makkah", tr="1", tg="4", note="Before")
        self.assertIn("Added 70 points to Makkah", done.get_data(as_text=True))
        self.assertEqual(self.client.get("/api/houses").get_json(), {"Makkah": 80, "Al-Aqsa": 0, "Madinah": 0})
        extra = self.client.get("/api/house-extra").get_json()
        self.assertEqual(sorted((e["coin"], e["n"], e["note"]) for e in extra), [("tg", 4, "Before"), ("tr", 1, "Before")])
        self.assertTrue(all(e["given_at"] > 1e12 for e in extra))
        self.assertEqual(tv.summary(self.query("SELECT id, json FROM classrooms"), extra={"Makkah": 70})["houses"][0],
                         {"house": "Makkah", "points": 80})
        page = self.client.get("/admin/house-points")
        self.client.post(f"/admin/house-points/{extra[0]['id']}/delete", data={"csrf": csrf_of(page)})
        self.assertEqual(len(self.client.get("/api/house-extra").get_json()), 1)
        # Coins counted on an earlier day are dated that day (noon, London); a day in the future is refused.
        late = self.form(self.client, "/admin/house-points", "/admin/house-points", house="Madinah", tg="2", day="2026-10-02")
        self.assertIn("Added 10 points to Madinah", late.get_data(as_text=True))
        dated = [e for e in self.client.get("/api/house-extra").get_json() if e["house"] == "Madinah"]
        self.assertEqual(datetime.fromtimestamp(dated[0]["given_at"] / 1000, auth.LONDON).strftime("%Y-%m-%d %H"), "2026-10-02 12")
        self.assertEqual(dated[0]["note"], "Coins counted")
        ahead = self.form(self.client, "/admin/house-points", "/admin/house-points", house="Madinah", tg="2", day="2999-01-01")
        self.assertEqual(ahead.status_code, 400)
        self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.assertEqual(teacher.get("/admin/house-points").status_code, 403)
        self.assertEqual(teacher.get("/api/house-extra").status_code, 403)

    def test_remove_quran_tab_coins_lists_them_first(self):
        t = 1760000000000
        amina = self.tv_student("Amina", house="Makkah", tg=[t, t + 1], tb=[t + 2, t + 3], tr=[t + 4],
                                hp=[{"t": t, "r": "hw"}, {"t": t + 4, "r": "bk"}],
                                tests=[{"d": "2026-10-13", "k": "m", "n": 3, "sc": 4, "ok": True, "tk": t + 2}])
        self.write(state=dict(self.state, S=[amina, self.tv_student("Bilal", house="Madinah", tr=[t])]))
        page = self.client.get("/admin/remove-quran-coins").get_data(as_text=True)
        self.assertIn("Amina", page)
        self.assertIn("Bilal", page)
        self.assertEqual(self.records()["yr3-a"]["revision"], 1)
        self.assertIn("Type DELETE", self.form(self.client, "/admin/remove-quran-coins", "/admin/remove-quran-coins",
                                               confirm="no").get_data(as_text=True))
        done = self.form(self.client, "/admin/remove-quran-coins", "/admin/remove-quran-coins", confirm="delete")
        self.assertIn("Removed 3 coins", done.get_data(as_text=True))
        state = json.loads(self.records()["yr3-a"]["json"])
        self.assertTrue(valid_state(state))
        a, b = state["S"]
        self.assertEqual((a["tg"], a["tb"], a["tr"], b["tr"]), ([t], [t + 2], [t + 4], []))
        self.assertIn("No student has house points from the Quran tab", self.client.get("/admin/remove-quran-coins").get_data(as_text=True))
        self.add_user("t@example.com")
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").get("/admin/remove-quran-coins").status_code, 403)

    def test_set_houses_matches_by_name_and_fills_in_ids_and_family(self):
        # On the list: Aayat Nm SE0001 SP457 Makkah, Musa Un SE0065, Adam Or SE0060 in Yr3-A and SE0265 in Yr3-B.
        yr3a = dict(self.state, S=[self.tv_student("Aayat Nm", sid="SE0005", house="Al-Aqsa", tg=[1]),
                                   self.tv_student("Musa Un", sid="SE0065"), self.tv_student("Adam Or", sid="SE0060"),
                                   self.tv_student("Saima K", sid="SE9999", house="Madinah")])
        self.write(name="Yr3-A-QM-Br Faeez", state=yr3a)
        self.write(name="Yr3B Ustadha Sophia", class_id="yr3-b", state=dict(self.state, S=[self.tv_student("Adam Or", sid="SE0060")]))
        page = self.client.get("/admin/set-houses").get_data(as_text=True)
        self.assertIn("Saima K (SE9999)", page)
        self.assertIn("Aayat Nm: SE0005 to SE0001", page)
        self.assertIn("Adam Or: SE0060 to SE0265", page)
        self.assertIn("SE0002 Ishaaq Ya", page)
        self.assertIn("Type HOUSES", self.form(self.client, "/admin/set-houses", "/admin/set-houses",
                                               confirm="no").get_data(as_text=True))
        done = self.form(self.client, "/admin/set-houses", "/admin/set-houses", confirm="houses")
        self.assertIn("Updated 4 students", done.get_data(as_text=True))
        records = self.records()
        a, b = (json.loads(records[k]["json"])["S"] for k in ("yr3-a", "yr3-b"))
        self.assertEqual([(s.get("sid"), s.get("fid", ""), s.get("house", "")) for s in a + b],
                         [("SE0001", "SP457", "Makkah"), ("SE0065", "SP290", "Al-Aqsa"), ("SE0060", "SP379", "Makkah"),
                          ("SE9999", "", "Madinah"), ("SE0265", "SP470", "Al-Aqsa")])
        self.assertEqual((a[0]["xid"], sorted(a[0]["ed"]), a[0]["tg"]),
                         (["SE0005"], ["fid", "gender", "house", "ist", "sid", "sy"], [1]))
        self.assertEqual((a[0]["sy"], a[0]["gender"], a[0]["ist"], b[0]["ist"]),
                         ("Year 1", "Female", "Yr1-A-AI-Sr Nasrin", "Yr3-B-AI-Sr Nusrat"))
        self.assertEqual(records["yr3-a"]["revision"], 2)
        self.assertTrue(valid_state(json.loads(records["yr3-a"]["json"])))
        teacher = self.sign_in("t@example.com", "teacher-password-1") if self.add_user("t@example.com") else None
        self.assertEqual(teacher.get("/admin/set-houses").status_code, 403)

    def test_import_memo_passes_clears_then_adds_the_sheet(self):
        sheet = ("Date,Student ID,Name,Target\n12/09/2026 10:36:25,SE0001,,L1 Target 2\n13/09/2026 10:00:00,se1,,L1 Target 2\n"
                 "20/09/2026 10:00:00,SE0001,,212. Surah Al-Mulk\n21/09/2026 10:00:00,SE0001,,Qalam page 2\n"
                 "26/09/2026 10:00:00,SE0001,,Fail\n27/09/2026 10:00:00,SE0001,,L1 Target 1\n"
                 "27/09/2026 10:00:00,,Musa Un,L1 Target 3\n27/09/2026 10:00:00,SE9999,Nobody Here,L1 Target 1\n"
                 "27/09/2026 10:00:00,,Musa Un,L9 Target 1\n28/09/2026 10:00:00,,Adam Or,L1 Target 5\n"
                 "25/10/2026 10:00:00,,Musa Un,L1 Target 2\n")
        url = "/admin/import-memo-passes"
        had = datetime(2026, 9, 1, 12, tzinfo=auth.LONDON).timestamp() * 1000
        quran_test = {"d": "2026-09-06", "k": "q", "n": 2, "sc": 3, "ok": True, "pq": 1}
        tests = [{"d": "2026-09-05", "k": "m", "n": 1, "sc": 4, "ok": True, "tk": had}, quran_test]
        self.write(state=dict(self.state, S=[self.tv_student("Aayat Nm", sid="SE0001", target=2, passes=1, log=[had], lt=[1], lx=[""],
                                                             fl=[had], tests=tests, tb=[had]),
                                             self.tv_student("Musa Un"),
                                             self.tv_student("Adam Or", sid="SE0060", target=4, passes=1, log=[had], lt=[3], lx=[""])]))
        self.write(name="Yr3-B", class_id="yr3-b", state=dict(self.state, S=[self.tv_student("Adam Or", sid="SE0265")]))
        hub = self.client.get("/admin/import").get_data(as_text=True)
        for text in ("New students", "/#upload-students", "Memorisation passes", "Qaidah/Tajweed level passes", "Choose CSV file"):
            self.assertIn(text, hub)
        self.assertIn("needs columns called Date, Student ID and Target",
                      self.form(self.client, url, url, csv=(io.BytesIO(b"Day,Who\n1,2\n"), "x.csv")).get_data(as_text=True))
        self.assertIn("isn&#39;t a date", self.form(self.client, url, url, csv=(io.BytesIO(b"Date,Student ID,Target\nsoon,A,B\n"),
                                                                                   "x.csv")).get_data(as_text=True))
        page = self.form(self.client, url, url, csv=(io.BytesIO(sheet.encode()), "passes.csv")).get_data(as_text=True)
        for text in ("No student on the site has this Student ID", "Two students share this name", "Not on the class&#39;s list",
                     "The same target is higher up the sheet", "possibly a revision target", "213. Surah-Al Qalam, page 2",
                     "Adam Or in Yr3-B (SE0265)", "2 passes, 1 fail and 1 Memo level test"):
            self.assertIn(text, page)
        # The page sends the sheet back with the confirmation.
        sent = re.search(r'<textarea name="sheet" hidden>(.*?)</textarea>', page, re.S).group(1)
        self.assertEqual(html.unescape(sent), sheet)
        self.assertIn("Type IMPORT", self.form(self.client, url, url, sheet=sheet, confirm="no").get_data(as_text=True))
        self.assertIn("Choose a CSV file first", self.form(self.client, url, url, confirm="import").get_data(as_text=True))
        for _ in range(2):  # a second go clears and adds the same again
            done = self.form(self.client, url, url, sheet=sheet, confirm="import", **{"pick-adamor": "yr3-b|0"}).get_data(as_text=True)
            self.assertIn("Imported the CSV into 2 classes", done)
            records = self.records()
            a, musa, other = json.loads(records["yr3-a"]["json"])["S"]
            self.assertTrue(valid_state(json.loads(records["yr3-a"]["json"])))
            self.assertEqual((a["lt"], a["lx"], a["passes"], a["target"], len(a["fl"])),
                             ([2, 3, 3, 1], ["", "212. Surah Al-Mulk", "213. Surah-Al Qalam, page 2", ""], 4, 3, 1))
            self.assertEqual(a["log"][0], datetime(2026, 9, 12, 10, 36, 25, tzinfo=auth.LONDON).timestamp() * 1000)
            # The Memo level test goes; its blue coin stays as a level completion.
            self.assertEqual((a["tests"], a["tb"], a["hp"]), ([quran_test], [had], [{"t": had, "r": "lv"}]))
            # A revision target passed later doesn't send them back in the book.
            self.assertEqual((musa["lt"], musa["target"], musa.get("tg", [])), ([3, 2], 4, []))
            self.assertEqual((other["log"], other["passes"], other["target"]), ([], 0, 4))
        b = json.loads(records["yr3-b"]["json"])["S"][0]
        self.assertEqual((b["lt"], b["target"]), ([5], 6))
        self.assertEqual(records["yr3-a"]["revision"], 3)

    def test_import_qaidah_passes_clears_then_adds_level_tests(self):
        sheet = ("Date,Student ID,Level,Score\n01/10/2026,SE0001,Level 1,\n08/10/2026,SE0001,Qaidah 3,4\n08/10/2026,se1,3,\n"
                 "15/10/2026,SE0001,Fail,\n10/10/2026,SE0001,Level 2,\n20/10/2026,SE0001,Level 21,\n")
        url = "/admin/import-qaidah-passes"
        memo_test = {"d": "2026-09-05", "k": "m", "n": 1, "sc": 4, "ok": True, "tk": 5}
        self.write(state=dict(self.state, S=[self.tv_student("Aayat Nm", sid="SE0001", qa=2, tb=[5, 6], tests=[
            memo_test, {"d": "2026-09-06", "k": "q", "n": 2, "sc": 3, "ok": True, "pq": 1, "tk": 6}]),
                                             self.tv_student("Musa Un", qa=5)]))
        page = self.form(self.client, url, url, csv=(io.BytesIO(sheet.encode()), "levels.csv")).get_data(as_text=True)
        for text in ("1 Quran level test already on the site", "Not a level from 1 to 20", "The same level is higher up the sheet",
                     "Lower than a level they had already passed", "Qaidah Level 3"):
            self.assertIn(text, page)
        done = self.form(self.client, url, url, sheet=sheet, confirm="import").get_data(as_text=True)
        self.assertIn("Imported the CSV into 1 class", done)
        state = json.loads(self.records()["yr3-a"]["json"])
        self.assertTrue(valid_state(state))
        a, musa = state["S"]
        self.assertEqual(a["tests"], [memo_test,
                                      {"d": "2026-10-01", "k": "q", "n": 1, "ok": True, "pq": 0},
                                      {"d": "2026-10-08", "k": "q", "n": 3, "ok": True, "pq": 1, "sc": 4},
                                      {"d": "2026-10-10", "k": "q", "n": 2, "ok": True, "pq": 3},
                                      {"d": "2026-10-15", "k": "q", "n": 3, "ok": False}])
        # The Qaidah level is the latest level passed; the cleared test's blue coin stays.
        self.assertEqual((a["qa"], a["tb"], a["hp"], musa["qa"]), (2, [5, 6], [{"t": 6, "r": "lv"}], 5))

    def test_activity_is_recorded_and_pruned(self):
        self.sign_in("someone@example.com", "wrong password")
        self.write()
        page = self.client.get("/admin/activity").get_data(as_text=True)
        for text in ("Signed in", "Wrong password", "someone@example.com", "Saved classroom"):
            self.assertIn(text, page)
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE audit_log SET at = at - ?", (15 * auth.DAY,))
            tx.execute("UPDATE classroom_history SET saved_at = saved_at - ?", (15 * auth.DAY,))
        auth.prune(self.app.config["DB"], 14)
        self.assertEqual(self.query("SELECT id FROM audit_log"), [])
        self.assertEqual(self.query("SELECT id FROM classroom_history"), [])
        self.assertEqual(list(self.records()), ["yr3-a"])

    def tv_student(self, name, **fields):
        return dict({"name": name, "target": 1, "passes": 0, "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [],
                     "ml": [], "q": {}}, **fields)

    def test_tv_dashboard_figures(self):
        now = datetime(2026, 10, 14, 12, tzinfo=auth.LONDON)  # a Wednesday
        ms = lambda *args: datetime(*args, tzinfo=auth.LONDON).timestamp() * 1000  # noqa: E731
        monday, last_month = ms(2026, 10, 12, 9), ms(2026, 9, 30, 9)
        rows = [{"id": "yr3-a", "json": json.dumps({"name": "Yr3-A-QM-Br Ahmed", "S": [
                    self.tv_student("Aisha Bint Khan", house="Makkah", log=[monday] * 3, ml=[monday], tr=[monday], tg=[1]),
                    self.tv_student("Yusuf Ali", house="Madinah", log=[monday] * 3 + [last_month], ml=[], tb=[1]),
                    self.tv_student("Old Book", year="Year 3", tr=[last_month], log=[last_month]),
                    self.tv_student("Bilal Noor", house="Madinah", bp=[{"t": monday, "b": "t"}]),
                    self.tv_student("Old Pass", bp=[{"t": last_month, "b": "q"}])]})},
                {"id": "yr8-b", "json": json.dumps({"name": "Yr8-B", "S": [
                    self.tv_student("Zara", year="Year 4", log=[monday]),  # her Safar year group wins over the class
                    self.tv_student("Omar Farooq", sy="Year 9", log=[monday] * 2)]})},
                {"id": "yr9-c", "json": json.dumps({"name": "Yr9-C", "S": [self.tv_student("Idle")]})}]
        data = tv.summary(rows, now)
        self.assertEqual(data["houses"], [{"house": "Makkah", "points": 55}, {"house": "Madinah", "points": 20},
                                          {"house": "Al-Aqsa", "points": 0}])
        self.assertEqual([(b["name"], b["class"]) for b in data["books"]], [("Aisha K.", "Yr3A Ustadh Ahmed"), ("Bilal N.", "Yr3A Ustadh Ahmed")])
        # The TV's stars are last week's (5 October), when nobody passed or earned a merit.
        self.assertEqual(data["stars"], [])
        self.assertEqual([(step["place"], step["n"], [w["name"] for w in step["who"]]) for step in data["memo"]],
                         [(1, 3, ["Aisha K.", "Yusuf A."]), (2, 2, ["Omar F."]), (3, 1, ["Zara"])])
        # Memo passes list every rank, not only the top three.
        self.assertEqual([step["place"] for step in tv.podium([{"name": str(n), "n": n} for n in range(5, 0, -1)], places=None)],
                         [1, 2, 3, 4, 5])
        self.assertEqual(data["qaidah"], [])
        self.assertEqual([(c["n"], [w["class"] for w in c["who"]]) for c in data["classes"]["memo"]["primary"]],
                         [(6, ["Yr3A Ustadh Ahmed"])])
        self.assertEqual([(c["n"], [w["class"] for w in c["who"]]) for c in data["classes"]["memo"]["secondary"]],
                         [(3, ["Yr8B"])])
        self.assertEqual(data["month"], "October")
        self.assertNotIn("Khan", json.dumps(data))

    def test_tv_podiums_and_stars(self):
        now = datetime(2026, 10, 22, 12, tzinfo=auth.LONDON)  # a Thursday in the fourth week
        ms = lambda *args: datetime(*args, tzinfo=auth.LONDON).timestamp() * 1000  # noqa: E731
        test = lambda d, ok=True, k="q": {"d": d, "k": k, "n": 2, "sc": 4, "ok": ok}  # noqa: E731
        rows = [{"id": "a", "json": json.dumps({"name": "Yr2-A-QM-Ustadh Faeez", "S": [
                    self.tv_student("Amina Rahman", house="Makkah", tests=[test("2026-10-01"), test("2026-10-08"),
                                                                           test("2026-09-30"), test("2026-10-09", ok=False)],
                                    log=[ms(2026, 9, 30, 9)] * 2, sw=["2026-10-12"]),
                    self.tv_student("Bilal Hussain", house="Al-Aqsa", tests=[test("2026-10-02")],
                                    log=[ms(2026, 10, 7, 9)] * 4, ml=[ms(2026, 10, 14, 9)], ist="Yr2-A-AI-Sr Maryam")]})},
                {"id": "b", "json": json.dumps({"name": "Yr9/10-B-QM-Br Fawaaz", "S": [
                    self.tv_student("Hana Ali", tests=[test("2026-10-03")], log=[ms(2026, 10, 21, 9)],
                                    ml=[ms(2026, 10, 13, 9)])]})},
                {"id": "c", "json": json.dumps({"name": "Juniors", "S": [
                    self.tv_student("Musa Khan", year="Year 1", tests=[test("2026-10-03")])]})}]
        # Teachers come from the classrooms (Quran and Islamic Studies) and from accounts, once each.
        data = tv.summary(rows, now, teachers=["Ustadh Ahmed", "fawaaz Rahman"])
        self.assertEqual([(step["place"], step["n"], [w["name"] for w in step["who"]]) for step in data["qaidah"]],
                         [(1, 2, ["Amina R."]), (2, 1, ["Bilal H.", "Hana A.", "Musa K."])])
        # Juniors is primary from its students' year groups; Yr9/10 is secondary from its name.
        self.assertEqual([(c["n"], [w["class"] for w in c["who"]]) for c in data["classes"]["qaidah"]["primary"]],
                         [(3, ["Yr2A Ustadh Faeez"]), (1, ["Juniors"])])
        self.assertEqual([(c["n"], [w["class"] for w in c["who"]]) for c in data["classes"]["qaidah"]["secondary"]],
                         [(1, ["Yr9/10B Ustadh Fawaaz"])])
        # Stars: one per class for last week (12 October), all on one page. The teacher's pick (Amina) wins over
        # Bilal's merit; elsewhere the most passes and merits win; classes with nothing that week have none.
        self.assertEqual([(s["label"], s["class"], s["name"]) for s in data["stars"]],
                         [("Week of 12 Oct", "Yr2A Ustadh Faeez", "Amina R."), ("Week of 12 Oct", "Yr9/10B Ustadh Fawaaz", "Hana A.")])
        self.assertEqual(data["teachers"], ["Ustadh Ahmed", "Ustadh Faeez", "Ustadh Fawaaz", "Ustadha Maryam"])
        self.assertEqual([tv.class_label(n) for n in ("Year 3A Quran - Ustadh Faeez", "Yr4-B-QM-Sr Zainab", "Yr5-C Idris",
                                                      "Reception-QM", "Hifz group", "Yr8-A-QM-Ustadha Kawthar/Sr",
                                                      "Yr8-B-QM-Kawthar/Sr")],
                         ["Yr3A Ustadh Faeez", "Yr4B Ustadha Zainab", "Yr5C Ustadh Idris", "Reception", "Hifz group",
                          "Yr8A Ustadha Kawthar", "Yr8B Ustadha Kawthar"])
        window = {"qaidah": ["16:00", "17:00"], "memo": ["23:30", "00:30"]}
        at = lambda h, m: datetime(2026, 10, 22, h, m, tzinfo=auth.LONDON)  # noqa: E731
        self.assertEqual(tv.focus(window, at(16, 30)), ["qaidah"])
        self.assertEqual(tv.focus(window, at(17, 0)), None)
        self.assertEqual(tv.focus(window, at(0, 10)), ["memo"])
        self.assertEqual(tv.focus({}, at(16, 30)), None)

    def test_presenter_shows_one_class_to_its_teachers(self):
        now = datetime(2026, 10, 22, 12, tzinfo=auth.LONDON)  # a Thursday in the fourth week
        ms = lambda *args: datetime(*args, tzinfo=auth.LONDON).timestamp() * 1000  # noqa: E731
        q = lambda *scores: {f"2026-10-{d:02d}": v for d, v in zip(range(1, 30), scores)}  # noqa: E731
        row = {"id": "yr3-a", "json": json.dumps({"name": "Yr3-A-QM-Br Ahmed", "S": [
            self.tv_student("Amina Rahman", house="Makkah", q=q(5, 5, 4, 5) | {"2026-09-30": 0},
                            log=[ms(2026, 10, 6, 9)] * 3, tg=[ms(2026, 10, 2, 9)], tb=[ms(2026, 9, 2, 9)], sw=["2026-10-05"]),
            self.tv_student("Bilal Hussain", house="Madinah", q=q(4, 4, 4), log=[ms(2026, 10, 14, 9)],
                            ml=[ms(2026, 10, 14, 9)] * 2, tr=[ms(2026, 10, 15, 9)], tb=[ms(2026, 10, 16, 9)]),
            self.tv_student("Hana Ali", house="Makkah", q=q(5, 5), tg=[ms(2026, 10, 3, 9)] * 2)]})}
        data = tv.class_summary(row, now)
        self.assertEqual(data["class"], "Yr3A Ustadh Ahmed")
        # Quran: the best average this month, from students with 3 or more scores (Hana has only 2).
        self.assertEqual([(s["n"], [w["name"] for w in s["who"]]) for s in data["quran"]], [(4.8, ["Amina R."]), (4.0, ["Bilal H."])])
        self.assertEqual([(s["n"], [w["name"] for w in s["who"]]) for s in data["memo"]], [(3, ["Amina R."]), (1, ["Bilal H."])])
        # House points earned this month only (Amina's blue token is from September; a book is a red token).
        self.assertEqual([(s["n"], [w["name"] for w in s["who"]]) for s in data["points"]],
                         [(70, ["Bilal H."]), (10, ["Hana A."]), (5, ["Amina R."])])
        self.assertEqual([(s["n"], [w["name"] for w in s["who"]]) for s in data["houses"]], [(70, ["Madinah"]), (15, ["Makkah"])])
        self.assertEqual([(s["n"], [w["name"] for w in s["who"]]) for s in data["merits"]], [(2, ["Bilal H."])])
        self.assertEqual([b["name"] for b in data["books"]], ["Bilal H."])
        self.assertEqual([(s["label"], s["name"]) for s in data["stars"]],
                         [("Week of 5 Oct", "Amina R."), ("Week of 12 Oct", "Bilal H.")])
        self.assertEqual(data["riders"], ["Amina R.", "Bilal H.", "Hana A."])
        self.assertNotIn("Rahman", json.dumps(data))
        # Teachers present their own classes; admins any class.
        self.write(class_id="yr3-a")
        self.write("Yr4-B", class_id="yr4-b")
        self.add_user("t@example.com", classes=["yr3-a"])
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        page = teacher.get("/presenter/yr3-a")
        self.assertEqual(page.status_code, 200)
        self.assertIn("script-src 'self' 'nonce-", page.headers["Content-Security-Policy"])
        self.assertEqual(teacher.get("/api/presenter/yr3-a").get_json()["class"], "Yr3A")
        self.assertEqual(teacher.get("/presenter/yr4-b").status_code, 404)
        self.assertEqual(teacher.get("/api/presenter/yr4-b").status_code, 404)
        self.assertEqual(self.client.get("/api/presenter/yr4-b").get_json()["class"], "Yr4B")
        self.assertEqual(self.client.get("/presenter/nope").status_code, 404)
        self.assertEqual(self.app.test_client().get("/api/presenter/yr3-a").status_code, 401)

    def test_tv_screen_links(self):
        student = self.tv_student("Aisha Khan", house="Makkah", tg=[1])
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get("/tv").status_code, 302)
        self.assertEqual(anonymous.get("/api/tv").status_code, 401)
        self.assertEqual(anonymous.get("/tv/not-a-real-link").status_code, 404)
        self.assertEqual(anonymous.get("/tv/not-a-real-link/data").status_code, 404)
        # Staff can preview it; only admins make links.
        self.add_user("t@example.com")
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        self.assertEqual(teacher.get("/tv").status_code, 200)
        self.assertEqual(teacher.get("/api/tv").get_json()["houses"][0], {"house": "Makkah", "points": 5})
        self.assertEqual(teacher.get("/admin/tv").status_code, 403)
        self.assertEqual(teacher.post("/admin/tv/new", data={"name": "Hall"}).status_code, 403)
        page = self.form(self.client, "/admin/tv", "/admin/tv/new", name="Admin area TV").get_data(as_text=True)
        link = re.search(r'value="http://localhost(/tv/[^"]+)"', page).group(1)
        self.assertEqual(self.query("SELECT COUNT(*) AS n FROM tv_screens WHERE token_hash = ?",
                                    (auth.token_hash(link[4:]),))[0]["n"], 1)
        response = anonymous.get(link)
        self.assertEqual(response.status_code, 200)
        self.assertIn("script-src 'self' 'nonce-", response.headers["Content-Security-Policy"])
        data = anonymous.get(link + "/data").get_json()
        self.assertEqual(data["houses"][0]["points"], 5)
        # The link opens nothing else.
        for path in ("/", "/api/classes", "/api/tv", "/admin/tv"):
            self.assertIn(anonymous.get(path).status_code, (302, 401))
        self.assertIn("Made a TV screen link", self.client.get("/admin/activity").get_data(as_text=True))
        screen_id = self.query("SELECT id FROM tv_screens")[0]["id"]
        # Settings: a daily Qaidah window; a half-filled window is refused.
        bad = self.form(self.client, "/admin/tv", f"/admin/tv/{screen_id}/settings", qaidah_from="16:00", qaidah_to="")
        self.assertEqual(bad.status_code, 400)
        page = self.form(self.client, "/admin/tv", f"/admin/tv/{screen_id}/settings", qaidah_from="00:00",
                         qaidah_to="23:59", memo_from="", memo_to="").get_data(as_text=True)
        self.assertIn("Qaidah only 00:00 to 23:59", page)
        self.assertEqual(teacher.post(f"/admin/tv/{screen_id}/settings", data={"qaidah_from": "01:00"}).status_code, 403)
        self.assertEqual(anonymous.get(link + "/data").get_json()["focus"], ["qaidah"])
        self.assertIsNone(teacher.get("/api/tv").get_json()["focus"])
        self.form(self.client, "/admin/tv", f"/admin/tv/{screen_id}/off")
        self.assertEqual(anonymous.get(link).status_code, 404)
        self.assertEqual(anonymous.get(link + "/data").status_code, 404)
        self.assertIn("No TV links yet", self.client.get("/admin/tv").get_data(as_text=True))

    def test_memo_board_figures(self):
        import memo_board
        now = datetime(2026, 10, 29, 12, tzinfo=auth.LONDON)  # a Thursday
        ms = lambda *args: datetime(*args, tzinfo=auth.LONDON).timestamp() * 1000  # noqa: E731
        fb = {"d": "2026-10-20", "k": "m", "n": 3, "sc": 2, "ok": False, "fb": "Revise the last page"}
        rows = [{"id": "yr3-a", "json": json.dumps({"name": "Yr3-A-QM-Br Ahmed", "L": [9, 9], "Q": 37, "S": [
                    self.tv_student("Amina Rahman", target=4, passes=3, log=[ms(2026, 10, 1, 9), ms(2026, 10, 13, 9), ms(2026, 10, 27, 9)],
                                    lt=[1, 2, 12], lx=["", "", ""], fl=[ms(2026, 10, 20, 9)],
                                    tests=[fb, {"d": "2026-10-27", "k": "m", "n": 12, "sc": 4, "ok": True}]),
                    self.tv_student("Bilal Hussain", target=20, passes=1, log=[ms(2026, 10, 14, 9)], lt=[19],
                                    lx=["197. Surah An-Naba'"]),
                    self.tv_student("Hana Ali", target=3, passes=1, log=[ms(2026, 9, 1, 9)], lt=[2], lx=[""],
                                    fl=[ms(2026, 10, 2, 9)] * 2),
                    self.tv_student("Musa Khan")]})},
                {"id": "yr8-b", "json": json.dumps({"name": "Yr8-B", "L": [9], "S": [
                    self.tv_student("Zara Noor", fl=[ms(2026, 10, 6, 9)])]})}]
        data = memo_board.board(rows, now)
        by = {s["name"]: s for s in data["students"]}
        a = by["Amina Rahman"]
        self.assertEqual((a["status"], a["weeks"], a["passes"], a["fails"], a["rate"]), ("on", 0, 3, 1, 75))
        self.assertEqual((a["level"], a["levelDate"], a["tested"], a["failDate"], a["target"]),
                         ("L2 T3", "2026-10-27", "2026-10-27", "2026-10-20", "L1 T4"))
        self.assertEqual((a["feedback"], a["feedbackDate"], a["feedbackOk"]), ("Revise the last page", "2026-10-20", False))
        # Two weeks since a surah pass: slowing down; the surah is written as chosen on Record.
        self.assertEqual((by["Bilal Hussain"]["status"], by["Bilal Hussain"]["weeks"], by["Bilal Hussain"]["level"],
                          by["Bilal Hussain"]["target"]), ("slow", 2, "197. Surah An-Naba'", "Item 199"))
        self.assertEqual((by["Hana Ali"]["status"], by["Hana Ali"]["weeks"], by["Hana Ali"]["rate"]), ("behind", 8, 33))
        self.assertEqual((by["Musa Khan"]["status"], by["Musa Khan"]["target"], by["Musa Khan"]["rate"]), ("none", None, None))
        self.assertEqual((by["Zara Noor"]["status"], by["Zara Noor"]["weeks"], by["Zara Noor"]["stage"], by["Zara Noor"]["year"]),
                         ("behind", None, "secondary", 8))
        self.assertEqual([c["label"] for c in data["classes"]], ["Yr3A Ustadh Ahmed", "Yr8B"])
        self.assertEqual(memo_board.target_label(56, [9, 9], 37), "Quran complete")
        # A link: short names, one stage, and without fails or feedback when made that way.
        short = memo_board.board(rows, now, short=True, fails=False, scope="primary")
        self.assertEqual([s["name"] for s in short["students"]], ["Amina R.", "Bilal H.", "Hana A.", "Musa K."])
        self.assertNotIn("fails", short["students"][0])
        self.assertNotIn("Revise", json.dumps(short))

    def test_memo_board_pages_and_links(self):
        student = self.tv_student("Aisha Khan", target=2, passes=1, log=[time.time() * 1000], lt=[1], lx=[""],
                                  tests=[{"d": "2026-10-27", "k": "m", "n": 1, "sc": 5, "ok": True, "fb": "Lovely"}])
        self.assertEqual(self.write(state=dict(self.state, S=[student])).status_code, 200)
        self.assertEqual(self.write("Yr4-B", class_id="yr4-b", state=dict(self.state, S=[self.tv_student("Omar Ali")])).status_code, 200)
        # Feedback has a length limit.
        long = dict(student, tests=[dict(student["tests"][0], fb="x" * 301)])
        self.assertEqual(self.write(state=dict(self.state, S=[long]), revision=1).status_code, 400)
        anonymous = self.app.test_client()
        self.assertEqual(anonymous.get("/board").status_code, 302)
        self.assertEqual(anonymous.get("/api/board").status_code, 401)
        self.assertEqual(anonymous.get("/board/not-a-link").status_code, 404)
        self.assertEqual(anonymous.get("/board/not-a-link/data").status_code, 404)
        page = self.client.get("/board")
        self.assertEqual(page.status_code, 200)
        self.assertIn("script-src 'self' 'nonce-", page.headers["Content-Security-Policy"])
        data = self.client.get("/api/board").get_json()
        self.assertTrue(data["admin"])
        self.assertEqual([s["name"] for s in data["students"]], ["Aisha Khan", "Omar Ali"])
        self.assertEqual(data["students"][0]["feedback"], "Lovely")
        # Teachers see their own classes and can't make links.
        self.add_user("t@example.com", classes=["yr4-b"])
        teacher = self.sign_in("t@example.com", "teacher-password-1")
        mine = teacher.get("/api/board").get_json()
        self.assertEqual(([s["name"] for s in mine["students"]], mine["admin"]), (["Omar Ali"], False))
        self.assertEqual(teacher.get("/admin/board").status_code, 403)
        page = self.form(self.client, "/admin/board", "/admin/board/new", name="Yr3 whiteboard", scope="class:yr3-a").get_data(as_text=True)
        link = re.search(r'value="http://localhost(/board/[^"]+)"', page).group(1)
        self.assertIn("without fails or feedback", page)
        self.assertEqual(anonymous.get(link).status_code, 200)
        shown = anonymous.get(link + "/data").get_json()
        self.assertEqual(([s["name"] for s in shown["students"]], shown["fails"], shown["title"], shown["fixedClass"]),
                         (["Aisha K."], False, "Yr3 whiteboard", True))
        self.assertNotIn("Lovely", json.dumps(shown))
        for path in ("/", "/api/classes", "/api/board", "/admin/board"):
            self.assertIn(anonymous.get(path).status_code, (302, 401))
        bad = self.form(self.client, "/admin/board", "/admin/board/new", name="x", scope="class:nope")
        self.assertEqual(bad.status_code, 400)
        self.form(self.client, "/admin/board", "/admin/board/new", name="Hall", scope="", fails="1")
        self.assertIn("Made a memo board link", self.client.get("/admin/activity").get_data(as_text=True))
        link_id = self.query("SELECT id FROM board_links WHERE name = 'Yr3 whiteboard'")[0]["id"]
        self.form(self.client, "/admin/board", f"/admin/board/{link_id}/off")
        self.assertEqual(anonymous.get(link).status_code, 404)
        self.assertEqual(anonymous.get(link + "/data").status_code, 404)

    def test_shared_styles_and_fonts(self):
        anonymous = self.app.test_client()
        css = anonymous.get("/auth.css").get_data(as_text=True)
        self.assertIn("--navy:#003b6b", css)  # the shared tokens, also inlined in the dashboard
        self.assertIn("details.user", css)
        font = anonymous.get("/fonts/atkinson-next-latin.woff2")
        self.assertEqual((font.status_code, font.mimetype), (200, "font/woff2"))
        self.assertIn("max-age", font.headers["Cache-Control"])
        self.assertEqual(anonymous.get("/fonts/../server.py").status_code, 404)
        self.assertEqual(anonymous.get("/fonts/other.woff2").status_code, 404)
        self.assertIn("font-src 'self'", anonymous.get("/login").headers["Content-Security-Policy"])
        # The opening on the sign-in page focuses the email field from /intro.js once it ends, not with autofocus,
        # so saved-password suggestions don't cover the animation. Other pages still run no scripts.
        login = anonymous.get("/login")
        self.assertIn("script-src 'self'", login.headers["Content-Security-Policy"])
        self.assertIn('src="/intro.js"', login.get_data(as_text=True))
        self.assertNotIn("autofocus", login.get_data(as_text=True))
        self.assertIn('<fieldset class="hold" disabled>', login.get_data(as_text=True))
        script = anonymous.get("/intro.js")
        self.assertEqual((script.status_code, script.mimetype), (200, "text/javascript"))
        self.assertNotIn("script-src", anonymous.get("/forgot").headers["Content-Security-Policy"])
        self.assertEqual(self.client.get("/admin/users").headers["Cache-Control"], "no-store")

    def test_help_or_requests(self):
        mailer = FakeMailer()
        self.app.config["MAILER"] = mailer
        self.add_user("maryam@example.com")
        teacher = self.sign_in("maryam@example.com", "teacher-password-1")
        self.assertEqual(teacher.get("/api/help").get_json()["items"], [])
        self.assertEqual(teacher.post("/api/help", json={"kind": "help", "body": "x"}).status_code, 403)
        self.assertEqual(teacher.post("/api/help", json={"kind": "other", "body": "x"}, headers=self.headers).status_code, 400)
        self.assertEqual(teacher.post("/api/help", json={"kind": "help", "body": "  \n "}, headers=self.headers).status_code, 400)
        sent = teacher.post("/api/help", json={"kind": "help", "body": "How do I <b>add</b> a pass?\n\n\n\nThanks"},
                            headers=self.headers).get_json()
        self.assertEqual(sent["items"][0]["body"], "How do I <b>add</b> a pass?\n\nThanks")
        teacher.post("/api/help", json={"kind": "request", "body": "A second TV please"}, headers=self.headers)
        # Teachers can't open FAQs; admins see both messages waiting, escaped.
        self.assertEqual(teacher.get("/admin/faqs").status_code, 403)
        page = self.client.get("/admin/faqs").get_data(as_text=True)
        self.assertIn("&lt;b&gt;add&lt;/b&gt;", page)
        self.assertIn("A second TV please", page)
        self.assertIn('href="/admin/faqs"', page)
        item = self.query("SELECT id FROM help_requests WHERE kind = 'help'")[0]["id"]
        replied = self.form(self.client, "/admin/faqs", f"/admin/faqs/{item}/reply", reply="Use Add a past pass.", send="email")
        self.assertIn("in the app and by email", replied.get_data(as_text=True))
        self.assertEqual(mailer.sent[-1][0], "maryam@example.com")
        self.assertIn("Use Add a past pass.", mailer.sent[-1][2])
        self.assertIn("&lt;b&gt;add&lt;/b&gt;", mailer.html[-1])
        mine = teacher.get("/api/help").get_json()
        self.assertEqual(mine["unread"], 1)
        self.assertEqual(next(i for i in mine["items"] if i["kind"] == "help")["reply"], "Use Add a past pass.")
        teacher.post("/api/help/seen", json={}, headers=self.headers)
        self.assertEqual(teacher.get("/api/help").get_json()["unread"], 0)
        # An app-only reply sends no email; another teacher sees none of this.
        other = self.query("SELECT id FROM help_requests WHERE kind = 'request'")[0]["id"]
        self.form(self.client, "/admin/faqs", f"/admin/faqs/{other}/reply", reply="Noted", send="app")
        self.assertEqual(len(mailer.sent), 1)
        self.add_user("idris@example.com")
        colleague = self.sign_in("idris@example.com", "teacher-password-1")
        self.assertEqual(colleague.get("/api/help").get_json()["items"], [])
        self.form(self.client, "/admin/faqs", f"/admin/faqs/{other}/delete")
        self.assertEqual(len(teacher.get("/api/help").get_json()["items"]), 1)

    def test_help_summary_every_sunday_at_noon(self):
        mailer = FakeMailer()
        self.app.config["MAILER"] = mailer
        london = auth.LONDON
        noon = datetime(2026, 10, 18, 12, tzinfo=london).timestamp()  # a Sunday
        self.assertEqual(helpdesk.last_summary_time(noon - 1).date().isoformat(), "2026-10-11")
        self.assertEqual(helpdesk.last_summary_time(noon).date().isoformat(), "2026-10-18")
        # Summer time ends on 25 October: still noon in London.
        self.assertEqual(helpdesk.last_summary_time(datetime(2026, 10, 25, 13, tzinfo=london).timestamp()).hour, 12)
        with self.app.test_request_context():
            # Nothing sent this week: the email says so.
            self.assertFalse(helpdesk.send_summary_if_due(noon - 60))
            self.assertTrue(helpdesk.send_summary_if_due(noon + 60))
            self.assertEqual(mailer.sent[-1][0], "mysafaradmin@gmail.com")
            self.assertIn("No help or requests were sent", mailer.sent[-1][2])
            # Only once a week, even when checked again.
            self.assertFalse(helpdesk.send_summary_if_due(noon + 120))
        # The next week's lists that week's messages, answered or not.
        self.add_user("maryam@example.com")
        teacher = self.sign_in("maryam@example.com", "teacher-password-1")
        teacher.post("/api/help", json={"kind": "request", "body": "More Qaidah books"}, headers=self.headers)
        week = datetime(2026, 10, 25, 12, tzinfo=london).timestamp()  # the clocks went back that morning
        with self.app.config["DB"].transaction() as tx:
            tx.execute("UPDATE help_requests SET created_at = ?", (int(week - 86400),))
        self.app.config["HELP_SUMMARY_TO"] = "office@example.com"
        with self.app.test_request_context():
            self.assertTrue(helpdesk.send_summary_if_due(week + 5))
            # Too late to send after the grace period.
            self.assertFalse(helpdesk.send_summary_if_due(week + 14 * 86400 + 2 * 86400))
        self.assertEqual(mailer.sent[-1][0], "office@example.com")
        self.assertIn("More Qaidah books", mailer.sent[-1][2])
        self.assertIn("Waiting for a reply", mailer.sent[-1][2])
        self.assertEqual(len(mailer.sent), 2)

    def test_password_help_links_to_the_school(self):
        anonymous = self.app.test_client()
        self.assertNotIn("mailto:", anonymous.get("/forgot").get_data(as_text=True))
        self.app.config["SUPPORT_EMAIL"] = "office@example.com"
        page = anonymous.get("/forgot").get_data(as_text=True)
        self.assertIn('href="mailto:office@example.com?subject=', page)

    # --- Security fixes ---------------------------------------------------------------

    def test_sign_in_never_follows_addresses_browsers_rewrite(self):
        # Browsers drop tabs and line breaks from an address, so "/\t/evil.example" would go to evil.example.
        for target in ("/\t/evil.example", "/\n/evil.example", "/\r/evil.example", "/ /evil.example", "/\x00/evil.example",
                       "/\x7f/evil.example", "/\\evil.example"):
            client = self.app.test_client()
            page = client.get("/login")
            response = client.post("/login", data={"csrf": csrf_of(page), "email": ADMIN[0], "password": ADMIN[1],
                                                   "next": target})
            self.assertEqual(response.headers["Location"], "/", repr(target))
        # Someone already signed in is sent on straight away, so the same check applies.
        self.assertEqual(self.client.get("/login?next=/%09/evil.example").headers["Location"], "/")
        self.assertEqual(self.client.get("/login?next=/admin/users%3Fa=1").headers["Location"], "/admin/users?a=1")
        self.assertEqual(self.app.test_client().head("/login").status_code, 200)

    def test_spreadsheet_keeps_formula_like_text_as_text(self):
        from io import BytesIO
        from openpyxl import load_workbook
        t = 1760000000000
        student = {"name": "=HYPERLINK(\"https://evil.example\",\"Open\")", "sid": "=1+1", "target": 1, "passes": 0,
                   "merits": 0, "warn": 0, "log": [], "lt": [], "lx": [], "ml": [], "q": {}}
        self.write(state=dict(self.state, S=[student]), name="=cmd|' /C calc'!A0")
        book = load_workbook(BytesIO(self.client.get("/admin/export.xlsx").data))
        row = [c for c in book["Students"][2]]
        self.assertEqual((row[0].value, row[1].value, row[2].value),
                         ("=cmd|' /C calc'!A0", "=1+1", "=HYPERLINK(\"https://evil.example\",\"Open\")"))
        self.assertTrue(all(c.data_type == "s" for c in row[:3]))

    def admin_mailer(self):
        mailer = FakeMailer()
        self.app.config["MAILER"] = mailer
        return mailer

    def code_of(self, mailer):
        return re.search(r"\b(\d{6})\b", mailer.sent[-1][2]).group(1)

    def test_admins_finish_signing_in_with_an_emailed_code(self):
        mailer = self.admin_mailer()
        client = self.sign_in(*ADMIN)
        # The password alone doesn't sign an admin in.
        self.assertEqual(client.last.status_code, 200)
        self.assertIn("Check your email", client.last.get_data(as_text=True))
        self.assertEqual(client.get("/api/classes").status_code, 401)
        self.assertEqual((mailer.sent[-1][0], mailer.tags[-1]), (ADMIN[0], "security"))
        code = self.code_of(mailer)
        # The code only works in the browser that typed the password.
        other = self.app.test_client()
        page = other.get("/login")
        self.assertEqual(other.post("/login/code", data={"csrf": csrf_of(page), "code": code}).status_code, 400)
        response = client.post("/login/code", data={"csrf": csrf_of(client.last), "code": code[:3] + " " + code[3:],
                                                     "next": "/admin/users"})
        self.assertEqual(response.headers["Location"], "/admin/users")
        self.assertEqual(client.get("/api/classes").status_code, 200)
        # It works once.
        page = client.get("/account")
        self.assertEqual(client.post("/login/code", data={"csrf": csrf_of(page), "code": code}).status_code, 400)
        # Teachers sign in with their password, as before.
        self.add_user("t@example.com")
        self.assertEqual(self.sign_in("t@example.com", "teacher-password-1").last.status_code, 302)

    def test_wrong_sign_in_codes_are_limited(self):
        mailer = self.admin_mailer()
        client = self.sign_in(*ADMIN)
        code = self.code_of(mailer)
        wrong = "000000" if code != "000000" else "111111"
        for attempt in range(auth.LOGIN_CODE_FAILS):
            response = client.post("/login/code", data={"csrf": csrf_of(client.last), "code": wrong})
            self.assertEqual(response.status_code, 401)
            client.last = response if attempt < auth.LOGIN_CODE_FAILS - 1 else client.get("/login")
        # After too many wrong codes, even the right one no longer works.
        page = client.get("/login")
        self.assertEqual(client.post("/login/code", data={"csrf": csrf_of(page), "code": code}).status_code, 400)
        self.assertEqual(client.get("/api/classes").status_code, 401)
        # Across new codes too: ten wrong codes in a quarter of an hour block new ones.
        for _ in range(2):
            client = self.sign_in(*ADMIN)
            for _ in range(auth.LOGIN_CODE_FAILS):
                client.post("/login/code", data={"csrf": csrf_of(client.get("/login")), "code": wrong})
        sent = len(mailer.sent)
        self.assertEqual(self.sign_in(*ADMIN).last.status_code, 429)
        self.assertEqual(len(mailer.sent), sent)

    def test_sign_in_code_that_cannot_be_emailed(self):
        mailer = self.admin_mailer()
        mailer.send = lambda *args, **kwargs: False
        client = self.sign_in(*ADMIN)
        self.assertEqual(client.last.status_code, 502)
        self.assertEqual(client.get("/api/classes").status_code, 401)
        self.assertEqual(self.query("SELECT COUNT(*) AS n FROM login_codes")[0]["n"], 0)

    def test_dashboard_scripts_carry_a_nonce(self):
        for path in ("/", "/tv", "/board"):
            response = self.client.get(path)
            nonce = re.search(r"'nonce-([^']+)'", response.headers["Content-Security-Policy"]).group(1)
            self.assertNotIn("unsafe-inline", response.headers["Content-Security-Policy"].split("style-src")[0])
            body = response.get_data(as_text=True)
            self.assertEqual(body.count("<script"), body.count(f'<script nonce="{nonce}"'), path)
            self.assertNotEqual(nonce, re.search(r"'nonce-([^']+)'", self.client.get(path).headers["Content-Security-Policy"]).group(1))
        self.assertIn("includeSubDomains", self.client.get("/").headers.get("Strict-Transport-Security", "includeSubDomains"))

    def test_errors_never_show_internal_details(self):
        app = create_app(dict(self.config, DATA_DIR=Path(self.temp.name) / "errors"))

        @app.get("/api/boom")
        @app.get("/boom")
        def boom():
            raise RuntimeError("secret detail /app/server.py SELECT * FROM users")

        with app.config["DB"].transaction() as tx:
            tx.execute("UPDATE users SET must_change = 0")
        client = self.sign_in(*ADMIN, app=app)
        with self.assertLogs(app.logger, "ERROR") as logs:
            for path in ("/boom", "/api/boom"):
                response = client.get(path)
                self.assertEqual(response.status_code, 500)
                self.assertNotIn("secret detail", response.get_data(as_text=True))
                self.assertNotIn("Traceback", response.get_data(as_text=True))
        self.assertIn("secret detail", logs.output[0])
        self.assertEqual(client.get("/api/boom").get_json(), {"error": server.ERRORS[500][1]})
        # API errors are JSON with their status; pages get the site's own error page.
        self.assertEqual(client.get("/api/nothing-here").get_json(), {"error": server.ERRORS[404][1]})
        self.assertIn("Page not found", client.get("/nothing-here").get_data(as_text=True))
        bad = client.put("/api/classes/yr3-a", data="{not json", content_type="application/json", headers=self.headers)
        self.assertEqual((bad.status_code, bad.get_json()), (400, {"error": server.ERRORS[400][1]}))
        # A message written for people is still shown.
        self.assertIn(auth.FORM_EXPIRED, html.unescape(client.post("/account/signout-others").get_data(as_text=True)))
        self.assertEqual(client.delete("/healthz").headers.get("Allow", "").count("GET"), 1)

    def test_profile_picture_limits(self):
        from PIL import Image
        # Too many pixels, even though the file is small.
        out = io.BytesIO()
        Image.new("L", (5000, 4000)).save(out, "PNG", optimize=True)
        self.assertLess(len(out.getvalue()), 2 * 1024 * 1024)
        self.assertEqual(self.upload_picture(self.client, out.getvalue()).status_code, 400)
        # Formats Pillow could open but we don't accept never reach their decoders.
        for fmt in ("TIFF", "BMP", "ICO"):
            self.assertEqual(self.upload_picture(self.client, self.picture(fmt=fmt)).status_code, 400, fmt)
        # A large JPEG photo is still fine.
        self.assertEqual(self.upload_picture(self.client, self.picture(size=(4000, 3000), fmt="JPEG")).status_code, 200)


class MailerTests(unittest.TestCase):
    def response(self, status=200):
        class Response:
            def __enter__(self):
                return self

            def __exit__(self, *args):
                return False
        Response.status = status
        return Response()

    def http_error(self, code):
        import io
        import urllib.error
        return urllib.error.HTTPError("https://api.resend.com/emails", code, "error", {}, io.BytesIO(b'{"message":"x"}'))

    def test_off_without_a_key_or_sender(self):
        from mailer import Mailer
        self.assertFalse(Mailer("", "Safar <a@example.org>").enabled)
        self.assertFalse(Mailer("re_key", "").send("t@example.com", "Hi", "Hello"))
        self.assertEqual(Mailer("re_key", "Safar Academy <dashboard@mysafar.org>").sender_address, "dashboard@mysafar.org")

    def test_sends_html_text_reply_to_and_tag(self):
        from mailer import Mailer
        with patch("urllib.request.urlopen", return_value=self.response()) as urlopen:
            ok = Mailer("re_key", "Safar <a@example.org>", reply_to="office@example.org").send(
                "t@example.com", "Hi", "Hello", html="<p>Hello</p>", tag="invite")
        self.assertTrue(ok)
        request = urlopen.call_args[0][0]
        body = json.loads(request.data)
        self.assertEqual(body, {"from": "Safar <a@example.org>", "to": ["t@example.com"], "subject": "Hi", "text": "Hello",
                                "html": "<p>Hello</p>", "reply_to": "office@example.org",
                                "tags": [{"name": "category", "value": "invite"}]})
        self.assertEqual(request.get_header("Authorization"), "Bearer re_key")
        self.assertTrue(request.get_header("Idempotency-key"))

    def test_test_mode_sends_everything_to_one_address(self):
        from mailer import Mailer
        with patch("urllib.request.urlopen", return_value=self.response()) as urlopen:
            Mailer("re_key", "a@example.org", redirect_to="me@example.org").send(
                "t@example.com", "Hi", "Hello", html="<td><!--TEST--><h1>Hi</h1></td>")
        body = json.loads(urlopen.call_args[0][0].data)
        self.assertEqual(body["to"], ["me@example.org"])
        self.assertIn("meant for t@example.com", body["text"])
        self.assertIn("meant for t@example.com]</p><h1>", body["html"])

    def test_retries_once_with_the_same_idempotency_key(self):
        from mailer import Mailer
        mailer = Mailer("re_key", "a@example.org")
        with patch("time.sleep"), patch("urllib.request.urlopen", side_effect=[self.http_error(503), self.response()]) as urlopen:
            self.assertTrue(mailer.send("t@example.com", "Hi", "Hello"))
        keys = {call[0][0].get_header("Idempotency-key") for call in urlopen.call_args_list}
        self.assertEqual((urlopen.call_count, len(keys)), (2, 1))
        with patch("time.sleep"), patch("urllib.request.urlopen", side_effect=self.http_error(422)) as urlopen, \
                patch("builtins.print"):
            self.assertFalse(mailer.send("t@example.com", "Hi", "Hello"))
        self.assertEqual(urlopen.call_count, 1)  # a rejected email isn't retried
        with patch("time.sleep"), patch("urllib.request.urlopen", side_effect=[self.http_error(429), self.http_error(429)]), \
                patch("builtins.print"):
            self.assertFalse(mailer.send("t@example.com", "Hi", "Hello"))

    def test_background_send_returns_straight_away(self):
        from mailer import Mailer
        import threading
        done = threading.Event()

        def deliver(payload):
            done.set()
            return True
        mailer = Mailer("re_key", "a@example.org")
        with patch.object(mailer, "_deliver", side_effect=deliver):
            self.assertTrue(mailer.send("t@example.com", "Hi", "Hello", background=True))
            self.assertTrue(done.wait(5))


class LogTests(unittest.TestCase):
    token = "Zx9-aB_cD3eF4gH5iJ6kL7mN8oP9qR0sT1uV2wX3yZ4"

    def test_access_log_hides_link_tokens(self):
        from datetime import timedelta
        from types import SimpleNamespace

        from gunicorn.config import Config
        conf = runpy.run_path(str(Path(__file__).with_name("gunicorn.conf.py")))
        logger = conf["logger_class"](Config())
        environ = {"REQUEST_METHOD": "POST", "RAW_URI": f"/reset/{self.token}?x=1", "PATH_INFO": f"/reset/{self.token}",
                   "QUERY_STRING": "x=1", "SERVER_PROTOCOL": "HTTP/1.1", "REMOTE_ADDR": "203.0.113.5"}
        request = SimpleNamespace(headers=[("REFERER", f"https://example.com/invite/{self.token}")])
        response = SimpleNamespace(status="302", response_length=0, sent=0, headers=[])
        atoms = logger.atoms(response, request, environ, timedelta(seconds=1))
        self.assertNotIn(self.token, json.dumps(atoms))
        self.assertEqual(atoms["r"], "POST /reset/[hidden]?x=1 HTTP/1.1")
        self.assertEqual(atoms["{referer}i"], "https://example.com/invite/[hidden]")

    def test_access_log_hides_tv_and_board_tokens(self):
        from datetime import timedelta
        from types import SimpleNamespace

        from gunicorn.config import Config
        conf = runpy.run_path(str(Path(__file__).with_name("gunicorn.conf.py")))
        logger = conf["logger_class"](Config())
        for path, shown in ((f"/tv/{self.token}/data", "/tv/[hidden]/data"), (f"/board/{self.token}", "/board/[hidden]"),
                            ("/admin/tv/3/off", "/admin/tv/3/off"), ("/admin/board/4/off", "/admin/board/4/off")):
            environ = {"REQUEST_METHOD": "GET", "RAW_URI": path, "PATH_INFO": path, "QUERY_STRING": "",
                       "SERVER_PROTOCOL": "HTTP/1.1", "REMOTE_ADDR": "203.0.113.5"}
            response = SimpleNamespace(status="200", response_length=0, sent=0, headers=[])
            atoms = logger.atoms(response, SimpleNamespace(headers=[]), environ, timedelta(seconds=1))
            self.assertNotIn(self.token, json.dumps(atoms))
            self.assertEqual(atoms["r"], f"GET {shown} HTTP/1.1")
        self.assertEqual(auth.hide_link_tokens(f"/tv/{self.token}"), "/tv/[hidden]")

    def test_error_log_hides_link_tokens(self):
        temp = tempfile.TemporaryDirectory()
        self.addCleanup(temp.cleanup)
        app = create_app({"TESTING": True, "DATABASE_URL": "", "DATA_DIR": temp.name,
                          "ADMIN_EMAIL": ADMIN[0], "DASHBOARD_PASSWORD": ADMIN[1]})
        with self.assertLogs(app.logger) as logs:
            app.logger.error("Exception on %s [POST]", f"/invite/{self.token}")
        self.assertEqual(logs.records[0].getMessage(), "Exception on /invite/[hidden] [POST]")


class SQLiteTest(DashboardTests, unittest.TestCase):
    def database_url(self):
        return ""

    def connect(self, config):
        Path(config["DATA_DIR"]).mkdir(parents=True, exist_ok=True)
        return sqlite3.connect(Path(config["DATA_DIR"]) / "dashboard.sqlite3")

    def execute(self, config, *statements):
        with closing(self.connect(config)) as db, db:
            for statement in statements:
                db.execute(statement)

    def table_exists(self, config, name):
        with closing(self.connect(config)) as db:
            return db.execute("SELECT 1 FROM sqlite_master WHERE type = 'table' AND name = ?", (name,)).fetchone() is not None

    def test_requires_database_on_railway(self):
        with patch.dict("os.environ", {"RAILWAY_ENVIRONMENT_ID": "test", "RAILWAY_VOLUME_MOUNT_PATH": ""}):
            with self.assertRaisesRegex(RuntimeError, "DATABASE_URL"):
                create_app(dict(self.config, TESTING=False))


@unittest.skipUnless(os.environ.get("TEST_DATABASE_URL"), "set TEST_DATABASE_URL to test PostgreSQL")
class PostgresTest(DashboardTests, unittest.TestCase):
    def database_url(self):
        url = os.environ["TEST_DATABASE_URL"]
        schema = "test_" + uuid.uuid4().hex
        with psycopg.connect(url, autocommit=True) as db:
            db.execute(f'CREATE SCHEMA "{schema}"')

        def drop():
            with psycopg.connect(url, autocommit=True) as db:
                db.execute(f'DROP SCHEMA "{schema}" CASCADE')
        self.addCleanup(drop)
        return make_conninfo(url, options=f"-c search_path={schema}")

    def execute(self, config, *statements):
        with psycopg.connect(config["DATABASE_URL"]) as db:
            for statement in statements:
                db.execute(statement)

    def table_exists(self, config, name):
        with psycopg.connect(config["DATABASE_URL"]) as db:
            return db.execute("SELECT to_regclass(%s)", (name,)).fetchone()[0] is not None

    def test_no_volume_needed_on_railway(self):
        with patch.dict("os.environ", {"RAILWAY_ENVIRONMENT_ID": "test", "RAILWAY_VOLUME_MOUNT_PATH": ""}):
            client = create_app(dict(self.config, TESTING=False)).test_client()
        self.assertEqual(client.get("/healthz").get_json(), {"status": "ok"})


if __name__ == "__main__":
    unittest.main()
